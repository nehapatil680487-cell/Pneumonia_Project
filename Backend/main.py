from fastapi import FastAPI, UploadFile, File, Form, HTTPException
from fastapi.responses import HTMLResponse, FileResponse, JSONResponse
from fastapi.staticfiles import StaticFiles
from database import get_connection

import tensorflow as tf
from tensorflow.keras.models import load_model

import numpy as np
import cv2
import os
import re
import uuid
import html
import traceback
import joblib
from datetime import datetime
from openpyxl import load_workbook

from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image,
    HRFlowable, KeepTogether, PageBreak
)
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT

# Corrected decision-fusion logic (see fusion.py)
from fusion import (
    combine_stage1, headline_for, result_class_for,
    LIKELY, POSSIBLE, INCONCLUSIVE, NEGATIVE,
)

# =========================
# APP SETUP
# =========================
app = FastAPI()

os.makedirs("static", exist_ok=True)
os.makedirs("reports", exist_ok=True)
app.mount("/static", StaticFiles(directory="static"), name="static")

# Folder for autofill.js (kept outside "static" because static/ is auto-cleaned)
os.makedirs("assets", exist_ok=True)
app.mount("/assets", StaticFiles(directory="assets"), name="assets")

IMG_SIZE = 224

# "Go to Home" buttons send the user back to the Streamlit landing page.
STREAMLIT_URL = "http://localhost:8501/"

# =========================
# MODEL OUTPUT POLARITY
# =========================
POSITIVE_CLASS_IS_PNEUMONIA_WHEN_PROB_HIGH = True

# The Oxygen ANN was trained on Celsius body temperature (e.g. 36.0, 36.1 ...).
# The application works in Fahrenheit, so the value is converted F -> C before prediction.
OXYGEN_MODEL_TEMP_UNIT = "C"

# Set True temporarily to print the exact inputs/outputs of the Oxygen ANN while testing.
OXYGEN_DEBUG = False

PDF_STORE = {}

SCREENING_DISCLAIMER = (
    "This is an AI-assisted screening aid, not a diagnosis. It supports, but does not replace, "
    "evaluation by a qualified doctor or radiologist. Please consult a doctor to confirm these results."
)


def cleanup_old_files(directory, max_age_seconds=3600):
    try:
        now = datetime.now().timestamp()
        for fname in os.listdir(directory):
            fpath = os.path.join(directory, fname)
            try:
                if os.path.isfile(fpath) and (now - os.path.getmtime(fpath)) > max_age_seconds:
                    os.remove(fpath)
            except OSError:
                pass
    except FileNotFoundError:
        pass


def cleanup_stale_pdf_store_entries():
    stale_ids = [rid for rid, path in PDF_STORE.items() if not os.path.exists(path)]
    for rid in stale_ids:
        PDF_STORE.pop(rid, None)


# =========================
# LOAD MODELS
# =========================
try:
    model = load_model("final_model.keras")
    try:
        _ = model(np.zeros((1, IMG_SIZE, IMG_SIZE, 3), dtype="float32"))
    except Exception as build_err:
        print("Model warm-up call failed (non-fatal):", build_err)
    print("Model loaded successfully")
except Exception as e:
    model = None
    print("Model load error:", e)

# --- Blood ANN (Stage 1 signal) ---
# If the model files have different names, change these three lines only.
BLOOD_MODEL_FILE = "blood_model (3).h5"
BLOOD_SCALER_FILE = "blood_scaler (1).pkl"
BLOOD_FEATURES_FILE = "blood_features.pkl"

try:
    blood_model = load_model(BLOOD_MODEL_FILE)
    blood_scaler = joblib.load(BLOOD_SCALER_FILE)
    try:
        # Saved by the training notebook: exact column order the model was trained on
        BLOOD_FEATURE_ORDER = list(joblib.load(BLOOD_FEATURES_FILE))
    except Exception:
        BLOOD_FEATURE_ORDER = ["Age", "Gender", "Hemoglobin", "Platelet_Count",
                               "White_Blood_Cells", "Red_Blood_Cells", "MCV", "MCH", "MCHC"]

    n_feat = len(BLOOD_FEATURE_ORDER)
    if getattr(blood_scaler, "n_features_in_", n_feat) != n_feat or blood_model.input_shape[-1] != n_feat:
        raise ValueError(
            f"Blood model/scaler/features files do not match "
            f"(features={n_feat}, scaler={getattr(blood_scaler, 'n_features_in_', '?')}, "
            f"model={blood_model.input_shape[-1]}). Copy all 3 files together."
        )

    _blood_mean = dict(zip(BLOOD_FEATURE_ORDER, blood_scaler.mean_))
    print("Blood ANN model loaded successfully. Features:", BLOOD_FEATURE_ORDER)
except Exception as e:
    blood_model = None
    blood_scaler = None
    BLOOD_FEATURE_ORDER = []
    _blood_mean = {}
    print("Blood ANN model load error:", e)

# --- Oxygen ANN (Stage 1 signal) ---
# Note: the files on disk have a space before the bracket, e.g. "oxygen_model (2).h5".
OXYGEN_MODEL_FILE = "oxygen_model (2).h5"
OXYGEN_SCALER_FILE = "oxygen_scaler (1).pkl"

try:
    oxygen_model = load_model(OXYGEN_MODEL_FILE)
    oxygen_scaler = joblib.load(OXYGEN_SCALER_FILE)
    # Training order was X = data[["pulse", "body temperature", "SpO2"]].
    # The scaler and network expect exactly this order.
    OXYGEN_FEATURE_ORDER = ["pulse_rate", "temperature", "spo2"]
    # Labels from the training notebook: 0 = Normal, 1 = Low Oxygen, 2 = Very Low Oxygen.
    OXYGEN_CLASS_LABELS = ["Normal", "Low oxygen", "Very low oxygen"]
    print("Oxygen ANN model loaded successfully")
    try:
        print("Oxygen scaler mean_ (expect [pulse, temp(C), spo2]):", oxygen_scaler.mean_)
    except Exception:
        pass
except Exception as e:
    oxygen_model = None
    oxygen_scaler = None
    print("Oxygen ANN model load error:", e)


# =========================
# PREPROCESS IMAGE
# =========================
def preprocess(file_bytes):
    img = cv2.imdecode(np.frombuffer(file_bytes, np.uint8), cv2.IMREAD_COLOR)
    if img is None:
        raise ValueError("Invalid image")
    img = cv2.resize(img, (IMG_SIZE, IMG_SIZE))
    img_norm = img.astype("float32") / 255.0
    return img, np.expand_dims(img_norm, axis=0)


def get_last_conv(model):
    last_name = None
    layer_owner = model

    def scan(m):
        nonlocal last_name, layer_owner
        for layer in m.layers:
            if isinstance(layer, tf.keras.layers.Conv2D):
                last_name = layer.name
                layer_owner = m
            elif isinstance(layer, tf.keras.Model):
                scan(layer)

    scan(model)
    return last_name, layer_owner


REFERENCE_RANGES = {
    "wbc":         {"label": "WBC",         "unit": "cells/mcL",  "low": 4000, "high": 11000},
    "rbc":         {"label": "RBC",         "unit": "mill/mcL",   "low": 4.2,  "high": 5.9},
    "platelet":    {"label": "Platelets",   "unit": "x10\u00b3/mcL", "low": 150, "high": 450},
    "hb":          {"label": "Hemoglobin",  "unit": "g/dL",       "low": 12.0, "high": 17.5},
    "temperature": {"label": "Temperature", "unit": "\u00b0F",    "low": 97.0, "high": 99.0},
    "spo2":        {"label": "SpO2",        "unit": "%",          "low": 95,   "high": 100},
    "pulse_rate":  {"label": "Pulse Rate",  "unit": "bpm",        "low": 60,   "high": 100},
}

# Lab rows that belong to the "Blood test" check in the Summary of Checks.
BLOOD_LAB_LABELS = ("WBC", "RBC", "Platelets", "Hemoglobin")


def evaluate_value(key, raw_value):
    ref = REFERENCE_RANGES[key]
    label = ref["label"]
    unit = ref["unit"]
    range_str = f"{ref['low']} - {ref['high']}"

    raw_str = "" if raw_value is None else str(raw_value).strip()
    if raw_str == "":
        return {"label": label, "value": "-", "unit": unit, "range": range_str, "flag": "N/A"}

    try:
        value = float(raw_str)
    except (ValueError, TypeError):
        return {"label": label, "value": "-", "unit": unit, "range": range_str, "flag": "N/A"}

    unit_warning = None
    if key == "platelet" and value > 1000:
        value = value / 1000.0
        unit_warning = "auto-converted from absolute count to x10\u00b3/mcL"
    elif key == "temperature" and 30 <= value <= 45:
        value = value * 9 / 5 + 32
        unit_warning = "auto-converted from \u00b0C to \u00b0F"

    if value < ref["low"]:
        flag = "Low"
    elif value > ref["high"]:
        flag = "High"
    else:
        flag = "Normal"

    value_str = f"{value:g}"
    if unit_warning:
        value_str += " *"

    return {"label": label, "value": value_str, "unit": unit,
            "range": range_str, "flag": flag, "unit_warning": unit_warning}


def evaluate_all_vitals(vitals):
    order = ["temperature", "spo2", "pulse_rate", "wbc", "rbc", "platelet", "hb"]
    return [evaluate_value(key, vitals.get(key, "")) for key in order]


def blood_values_out_of_range(lab_rows):
    """
    Blood values (WBC, RBC, Platelets, Hemoglobin) that fall outside the reference range,
    e.g. ["WBC high", "RBC low", "Hemoglobin low"]. This is a rule-based range check and is
    independent of the Blood ANN's pneumonia probability.
    """
    return [f"{r['label']} {r['flag'].lower()}"
            for r in lab_rows
            if r["label"] in BLOOD_LAB_LABELS and r["flag"] in ("High", "Low")]


# =========================================================================
# THREE-LEVEL STATUS GRADING (Normal / Borderline / Abnormal)
#   Normal     = within range
#   Borderline = slightly outside range
#   Abnormal   = clearly outside range
# Used for the "Oxygen & pulse" row (SpO2 + pulse + temperature; the worst value wins)
# of the Summary of Checks.
# General screening ranges, not diagnostic cut-offs: confirm against clinical guidelines.
# =========================================================================
GRADE_ORDER = {"Normal": 0, "Borderline": 1, "Abnormal": 2}
GRADE_PILL = {"Normal": "neg", "Borderline": "warn", "Abnormal": "pos"}

# SpO2 thresholds depend on the patient's situation.
#   value >= normal_min      -> Normal
#   value >= borderline_min  -> Borderline
#   otherwise                -> Abnormal
SPO2_SITUATIONS = {
    "adult":     {"label": "healthy adult",                     "normal_min": 95, "borderline_min": 91},
    "child":     {"label": "child / infant",                    "normal_min": 95, "borderline_min": 92},
    "pneumonia": {"label": "pneumonia suspected / symptomatic", "normal_min": 95, "borderline_min": 92},
}


def pick_spo2_situation(age_raw, is_positive):
    """Choose which SpO2 threshold row applies, from age and the screening result."""
    try:
        if float(age_raw) < 18:
            return "child"
    except (ValueError, TypeError):
        pass
    if is_positive:
        return "pneumonia"
    return "adult"


def grade_spo2(value, situation_key):
    s = SPO2_SITUATIONS[situation_key]
    if value >= s["normal_min"]:
        return "Normal"
    if value >= s["borderline_min"]:
        return "Borderline"
    return "Abnormal"


def grade_pulse(value):
    if 60 <= value <= 100:
        return "Normal"
    if 50 <= value < 60 or 100 < value <= 110:
        return "Borderline"
    return "Abnormal"


def grade_temperature_f(value):
    # 36.1-37.2 C = 97.0-99.0 F normal; 37.3-38.0 C = 99.1-100.4 F borderline; > 38.0 C abnormal
    if value > 100.4:
        return "Abnormal"
    if value > 99.0:
        return "Borderline"
    if value >= 97.0:
        return "Normal"
    if value >= 95.0:          # mildly low temperature
        return "Borderline"
    return "Abnormal"          # clearly low (hypothermia range)


def _lab_value(lab_rows, label):
    row = next((r for r in lab_rows if r["label"] == label), None)
    if row is None or row["flag"] == "N/A":
        return None
    try:
        return float(str(row["value"]).rstrip(" *"))
    except (ValueError, TypeError):
        return None


def grade_oxygen_group(lab_rows, situation_key):
    """Grade SpO2, pulse and temperature; the row's final label is the WORST of the available ones."""
    spo2 = _lab_value(lab_rows, "SpO2")
    pulse = _lab_value(lab_rows, "Pulse Rate")
    temp = _lab_value(lab_rows, "Temperature")

    parts, grades = [], []
    spo2_grade = None
    if spo2 is not None:
        spo2_grade = grade_spo2(spo2, situation_key)
        grades.append(spo2_grade)
        parts.append(f"SpO2 {spo2:g}% - {spo2_grade}")
    if pulse is not None:
        g = grade_pulse(pulse)
        grades.append(g)
        parts.append(f"Pulse {pulse:g} bpm - {g}")
    if temp is not None:
        g = grade_temperature_f(temp)
        grades.append(g)
        parts.append(f"Temperature {temp:g} \u00b0F - {g}")

    s = SPO2_SITUATIONS[situation_key]
    if not grades:
        return {"available": False, "label": None, "spo2_grade": None,
                "situation_key": situation_key, "situation_label": s["label"], "breakdown": ""}

    worst = max(grades, key=lambda g: GRADE_ORDER[g])
    return {"available": True, "label": worst, "spo2_grade": spo2_grade,
            "situation_key": situation_key, "situation_label": s["label"],
            "breakdown": "; ".join(parts)}


# =========================
# STAGE 1: MULTIMODAL PNEUMONIA DETECTION
# (the fusion / decision rule itself lives in fusion.py -> combine_stage1)
# =========================
def _blood_gender_value(gender_raw):
    g = str(gender_raw or "").strip().lower()
    if g in ("male", "m"):
        return 1.0
    if g in ("female", "f"):
        return 0.0
    return _blood_mean.get("Gender", 0.5)      # "Other"/unknown -> training average


def predict_blood_ann(age_raw, hb_raw, platelet_raw, wbc_raw, rbc_raw, gender_raw=""):
    if blood_model is None or blood_scaler is None:
        return {"available": False, "reason": "Blood ANN model unavailable."}

    # Form fields -> (display name, raw value). Only the ones the model actually uses are required.
    form_inputs = {
        "Hemoglobin":        ("Hemoglobin", hb_raw),
        "Platelet_Count":    ("Platelets", platelet_raw),
        "White_Blood_Cells": ("WBC", wbc_raw),
        "Red_Blood_Cells":   ("RBC", rbc_raw),
    }
    needed = [f for f in BLOOD_FEATURE_ORDER if f in form_inputs]

    missing = [form_inputs[f][0] for f in needed if str(form_inputs[f][1]).strip() == ""]
    if missing:
        return {"available": False,
                "reason": f"Missing required value(s) for Blood ANN: {', '.join(missing)}."}

    try:
        parsed = {f: float(form_inputs[f][1]) for f in needed}
    except (ValueError, TypeError):
        return {"available": False, "reason": "Blood ANN inputs could not be parsed as numbers."}

    if "Platelet_Count" in parsed:
        parsed["Platelet_Count"] *= 1000.0     # form is x10^3/uL, model was trained on absolute count

    try:
        age = float(age_raw)
    except (ValueError, TypeError):
        age = _blood_mean.get("Age", 40.0)

    gender_val = _blood_gender_value(gender_raw)

    try:
        row = []
        for f in BLOOD_FEATURE_ORDER:
            if f == "Age":
                v = age
            elif f == "Gender":
                v = gender_val
            elif f in parsed:
                v = parsed[f]
            else:
                v = _blood_mean.get(f)          # MCV / MCH / MCHC are not on the form -> training average
            if v is None:
                return {"available": False, "reason": f"Blood ANN feature '{f}' has no value."}
            row.append(v)

        X = np.array([row], dtype="float32")
        X_scaled = blood_scaler.transform(X)
        out = np.asarray(blood_model.predict(X_scaled, verbose=0)).ravel()
        prob = float(out[-1])                   # works for sigmoid (1 value) and softmax (2 values: class 1)
        return {"available": True, "prob": prob, "is_positive": prob > 0.5, "partial_profile": True}
    except Exception as e:
        print("Blood ANN prediction error:", e)
        return {"available": False, "reason": "Blood ANN prediction failed."}


def predict_oxygen_ann(spo2_raw, pulse_rate_raw, temperature_raw):
    if oxygen_model is None or oxygen_scaler is None:
        return {"available": False, "reason": "Oxygen ANN model unavailable."}

    required = {"SpO2": spo2_raw, "Pulse Rate": pulse_rate_raw, "Temperature": temperature_raw}
    missing = [k for k, v in required.items() if str(v).strip() == ""]
    if missing:
        return {"available": False,
                "reason": f"Missing required value(s) for Oxygen ANN: {', '.join(missing)}."}

    try:
        spo2 = float(spo2_raw)
        pulse_rate = float(pulse_rate_raw)
        temperature = float(temperature_raw)   # already cleaned to degrees F
    except (ValueError, TypeError):
        return {"available": False, "reason": "Oxygen ANN inputs could not be parsed as numbers."}

    # Convert F -> C because the model was trained on Celsius.
    if OXYGEN_MODEL_TEMP_UNIT == "C":
        temperature = (temperature - 32.0) * 5.0 / 9.0

    try:
        # Feature order must match training: [pulse, body temperature, SpO2].
        X = np.array([[pulse_rate, temperature, spo2]], dtype="float32")
        X_scaled = oxygen_scaler.transform(X)
        probs = oxygen_model.predict(X_scaled, verbose=0).squeeze()
        probs = [float(p) for p in probs]
        class_idx = int(np.argmax(probs))
        p_pos = 1.0 - probs[0]            # class 0 = Normal
        if OXYGEN_DEBUG:
            print(f"[oxygen] X(pulse,tempC,spo2)={X.tolist()} scaled={X_scaled.tolist()} "
                  f"probs={[round(p, 4) for p in probs]} class={OXYGEN_CLASS_LABELS[class_idx]}")
        return {
            "available": True,
            "class_label": OXYGEN_CLASS_LABELS[class_idx],
            "class_probs": probs,
            "p_pos": p_pos,
            "is_positive": p_pos > 0.5,
        }
    except Exception as e:
        print("Oxygen ANN prediction error:", e)
        return {"available": False, "reason": "Oxygen ANN prediction failed."}


# ---------- Patient-friendly wording ----------
def needs_urgent_review(is_positive, oxygen_grade):
    """
    Safety rule: when the combined result is negative but the Oxygen & pulse grade is Abnormal,
    the report must not present an unqualified "no signs" message.
    """
    return bool(not is_positive and oxygen_grade
                and oxygen_grade.get("available") and oxygen_grade.get("label") == "Abnormal")


# (row name, positive text, negative text, note)
FRIENDLY_SIGNALS = {
    "X-ray CNN":    ("Chest X-ray", "Signs of pneumonia", "Looks normal", "AI review of the chest image"),
    "Blood ANN":    ("Blood test", "Needs review", "Normal", "Based on the blood values entered"),
    "Oxygen ANN":   ("Oxygen & pulse", "Abnormal", "Normal", "Based on oxygen, pulse and temperature"),
    "Final Result": ("Overall screening", "Pneumonia suspected", "No signs found",
                     "Combined result of the checks above"),
}


def friendly_signal(c, oxygen_grade=None, blood_abnormal=None, urgent_override=False):
    name, pos, neg, note = FRIENDLY_SIGNALS.get(c["name"], (c["name"], c["result"], c["result"], ""))

    # Oxygen & pulse: rule-based 3-label grade (worst of SpO2 / pulse / temperature),
    # with SpO2 thresholds chosen by the patient's situation.
    if c["name"] == "Oxygen ANN":
        if oxygen_grade and oxygen_grade.get("available"):
            g = oxygen_grade["label"]
            return name, g, GRADE_PILL[g], note
        return name, "Not tested", "na", ""

    if not c["available"]:
        return name, "Not tested", "na", ""

    # Blood test: Blood ANN found no pneumonia pattern, but values are outside the reference range.
    if c["name"] == "Blood ANN" and not c["is_positive"] and blood_abnormal:
        return name, "Abnormal values", "warn", "No pneumonia pattern, but some values are out of range"

    # Overall screening: wording follows the fused verdict (LIKELY / POSSIBLE / INCONCLUSIVE / NEGATIVE).
    if c["name"] == "Final Result":
        v = c.get("verdict")
        if v == LIKELY:
            return name, "Pneumonia suspected", "pos", note
        if v == POSSIBLE:
            if c.get("driver") == "blood":
                return name, "Needs doctor review", "warn", "X-ray normal, but blood results need review"
            return name, "Possible pneumonia", "warn", "Some findings need a doctor's review"
        if v == INCONCLUSIVE:
            return name, "Inconclusive", "warn", "Checks disagree; doctor review required"
        if urgent_override:
            return name, "Urgent review advised", "warn", "Abnormal vital signs despite no pneumonia pattern"
        return name, "No signs found", "neg", note

    return name, (pos if c["is_positive"] else neg), ("pos" if c["is_positive"] else "neg"), note


def technical_signal(c, blood_abnormal=None, oxygen_grade=None):
    """
    Wording for the 'Clinical details (for doctors)' section (web page and PDF appendix).
    Returns (pill_class, pill_text, detail).
    """
    if not c["available"]:
        return "na", "Not assessed", c["detail"]

    if c["name"] == "Blood ANN":
        if c["is_positive"]:
            return "pos", "Pneumonia pattern", c["detail"]
        if blood_abnormal:
            detail = f"{c['detail']} | Out of reference range: {', '.join(blood_abnormal)}"
            return "warn", "Abnormal values", detail
        return "neg", "No pneumonia pattern", c["detail"]

    if c["name"] == "Oxygen ANN":
        # Same rule-based grade (SpO2 / pulse / temperature, worst value) as the patient summary,
        # so both sections always agree. The ANN output is shown as supporting detail.
        if oxygen_grade and oxygen_grade.get("available"):
            g = oxygen_grade["label"]
            return GRADE_PILL[g], g, f"ANN: {c['detail']} | Rule-based grade: {g}"
        return ("pos", "Abnormal", c["detail"]) if c["is_positive"] else ("neg", "Normal", c["detail"])

    if c["name"] == "Final Result":
        pill = {LIKELY: "pos", POSSIBLE: "warn", INCONCLUSIVE: "warn"}.get(c.get("verdict"), "neg")
        return pill, c["result"], c["detail"]

    return ("pos" if c["is_positive"] else "neg"), c["result"], c["detail"]


# =========================
# STAGE 2: CAP SEVERITY (NSCAP / SCAP)
# =========================
SEVERITY_INFO = {
    "NSCAP": {
        "full_name": "Non-Severe Community-Acquired Pneumonia",
        "description": (
            "Fewer than three ATS/IDSA minor severity criteria were met among the "
            "criteria that could be evaluated from the entered values. Non-severe CAP is "
            "typically manageable with standard outpatient or general-ward care, though "
            "clinical correlation is still required."
        ),
    },
    "SCAP": {
        "full_name": "Severe Community-Acquired Pneumonia",
        "description": (
            "Three or more ATS/IDSA minor severity criteria were met. Severe CAP is "
            "associated with a higher-risk clinical course and often warrants closer "
            "monitoring or escalated (e.g. ICU-level) care."
        ),
    },
}

# Patient-facing severity wording (plain language, no medical abbreviations needed).
PATIENT_SEVERITY = {
    "mild": {
        "title": "Mild to Moderate",
        "tag": "NSCAP",
        "text": ("Severity appears mild to moderate based on the values entered. No serious warning "
                 "signs were found. This kind of pneumonia can often be treated with medicines and "
                 "rest, but please have a doctor confirm this."),
    },
    "severe": {
        "title": "Severe \u2013 See a Doctor Urgently",
        "tag": "SCAP",
        "text": ("Serious warning signs were found. Please go to a hospital or emergency care "
                 "as soon as possible."),
    },
}

SPO2_HYPOXEMIA_CUTOFF = 90
ATS_IDSA_MINOR_CRITERIA = [
    ("Leukopenia (WBC < 4000)",          "WBC",         lambda v: v < 4000, "cells/mcL"),
    ("Thrombocytopenia (Plt < 100)",     "Platelets",   lambda v: v < 100,  "x10\u00b3/mcL"),
    ("Hypothermia (Temp < 96.8\u00b0F)", "Temperature", lambda v: v < 96.8, "\u00b0F"),
    (f"Hypoxemia (SpO2 < {SPO2_HYPOXEMIA_CUTOFF}, proxy)", "SpO2",
     lambda v: v < SPO2_HYPOXEMIA_CUTOFF, "%"),
]


def heatmap_hot_fraction(heatmap):
    """Kept for reference only. Grad-CAM attention area is NOT used in severity scoring."""
    h = cv2.resize(heatmap.astype("float32"), (IMG_SIZE, IMG_SIZE), interpolation=cv2.INTER_CUBIC)
    if float(h.max()) <= 1e-6:
        return 0.0
    return float((h >= 0.6 * h.max()).sum()) / h.size


def predict_severity(lab_rows, spo2_abnormal=None):
    """
    spo2_abnormal: True/False when the situation-based SpO2 grade is known
    (Abnormal = clearly low for this patient's situation); None falls back to the fixed 90% cut-off.
    """
    criteria = []
    met = 0
    evaluated = 0

    for name, label, test, unit in ATS_IDSA_MINOR_CRITERIA:
        v = _lab_value(lab_rows, label)
        if v is None:
            criteria.append({"name": name, "status": "Not provided", "detail": "not entered", "met": False})
            continue
        evaluated += 1
        is_met = bool(test(v))
        met += int(is_met)
        criteria.append({"name": name, "status": "Met" if is_met else "Not met",
                         "detail": f"{v:g} {unit}", "met": is_met})

    if evaluated == 0:
        return {"available": False,
                "reason": "No severity-evaluable value (WBC, Platelets, Temperature, SpO2) was entered."}

    missing = len(criteria) - evaluated
    is_severe = met >= 3
    provisional = (not is_severe) and (met + missing >= 3)
    label = "SCAP" if is_severe else "NSCAP"
    meta = SEVERITY_INFO[label]

    # Safety rule: a clearly low oxygen level always triggers the urgent message,
    # regardless of the criteria count.
    spo2_val = _lab_value(lab_rows, "SpO2")
    if spo2_abnormal is None:
        urgent_override = (spo2_val is not None and spo2_val < SPO2_HYPOXEMIA_CUTOFF)
    else:
        urgent_override = bool(spo2_abnormal)
    level = "severe" if (is_severe or urgent_override) else "mild"

    description = meta["description"]
    if provisional:
        description += (f" Note: {missing} criterion/criteria could not be evaluated because the "
                        f"value was not entered, so this result is provisional.")

    return {
        "available": True,
        "label": label,
        "level": level,
        "urgent_override": urgent_override,
        "patient": PATIENT_SEVERITY[level],
        "full_name": meta["full_name"],
        "description": description,
        "criteria": criteria,
        "met": met,
        "evaluated": evaluated,
        "total": len(criteria),
        "provisional": provisional,
    }


# =========================
# GRAD-CAM
# =========================
def gradcam(model, img_array, layer_name, layer_owner):
    if model is None or layer_name is None:
        return np.zeros((IMG_SIZE, IMG_SIZE), dtype="float32")

    try:
        if not getattr(model, "inputs", None):
            _ = model(img_array)

        grad_model = tf.keras.models.Model(
            inputs=model.inputs,
            outputs=[layer_owner.get_layer(layer_name).output, model.output]
        )

        with tf.GradientTape() as tape:
            conv_out, preds = grad_model(img_array)
            loss = tf.reduce_max(preds, axis=1)

        grads = tape.gradient(loss, conv_out)
        if grads is None:
            return np.zeros((IMG_SIZE, IMG_SIZE), dtype="float32")

        pooled = tf.reduce_mean(grads, axis=(0, 1, 2))
        conv_out = conv_out[0]
        heatmap = conv_out @ pooled[..., tf.newaxis]
        heatmap = tf.squeeze(heatmap)
        heatmap = tf.maximum(heatmap, 0)
        heatmap = heatmap / (tf.reduce_max(heatmap) + 1e-8)
        return heatmap.numpy()

    except Exception as e:
        print("Grad-CAM error:", e)
        return np.zeros((IMG_SIZE, IMG_SIZE), dtype="float32")


def make_overlay(img, heatmap):
    heatmap = cv2.resize(heatmap, (IMG_SIZE, IMG_SIZE))
    heatmap = np.uint8(255 * heatmap)
    heatmap = cv2.applyColorMap(heatmap, cv2.COLORMAP_JET)
    img = cv2.resize(img, (IMG_SIZE, IMG_SIZE))
    return cv2.addWeighted(img, 0.6, heatmap, 0.4, 0)


def get_bbox_from_heatmap(heatmap, threshold_ratio=0.6, min_area_ratio=0.008):
    h = cv2.resize(heatmap.astype("float32"), (IMG_SIZE, IMG_SIZE))
    if float(h.max()) <= 1e-6:
        return None

    mask = (h >= threshold_ratio * h.max()).astype("uint8") * 255
    kernel = np.ones((5, 5), np.uint8)
    mask = cv2.morphologyEx(mask, cv2.MORPH_CLOSE, kernel)
    mask = cv2.morphologyEx(mask, cv2.MORPH_OPEN, kernel)

    contours, _ = cv2.findContours(mask, cv2.RETR_EXTERNAL, cv2.CHAIN_APPROX_SIMPLE)
    if not contours:
        return None

    largest = max(contours, key=cv2.contourArea)
    if cv2.contourArea(largest) < min_area_ratio * IMG_SIZE * IMG_SIZE:
        return None

    x, y, w, box_h = cv2.boundingRect(largest)
    pad = 4
    return (max(0, x - pad), max(0, y - pad),
            min(IMG_SIZE, x + w + pad), min(IMG_SIZE, y + box_h + pad))


def draw_bbox(img, bbox, color=(0, 255, 255), thickness=2):
    out = img.copy()
    if bbox is not None:
        x1, y1, x2, y2 = bbox
        cv2.rectangle(out, (x1, y1), (x2, y2), color, thickness)
    return out


def _add_caption(image, text):
    bar_h = 40
    canvas = np.zeros((image.shape[0] + bar_h, image.shape[1], 3), dtype=np.uint8)
    canvas[bar_h:, :, :] = image
    cv2.putText(canvas, text, (10, 27), cv2.FONT_HERSHEY_SIMPLEX, 0.65,
                (255, 255, 255), 2, cv2.LINE_AA)
    return canvas


def make_comparison_image(img, heatmap, bbox):
    img_r = cv2.resize(img, (IMG_SIZE, IMG_SIZE))
    overlay = make_overlay(img, heatmap)
    left = _add_caption(draw_bbox(img_r, bbox), "Original Image")
    right = _add_caption(draw_bbox(overlay, bbox), "Image + Grad-CAM")
    gap = np.zeros((left.shape[0], 10, 3), dtype=np.uint8)
    return np.hstack([left, gap, right])


def interpret_heatmap(heatmap, xray_positive, bbox=None):
    """
    xray_positive must be the X-RAY CNN's own call (not the fused result), so the Grad-CAM text
    always describes what the image model actually predicted.
    """
    try:
        h = cv2.resize(heatmap.astype("float32"), (IMG_SIZE, IMG_SIZE), interpolation=cv2.INTER_CUBIC)

        if float(h.max()) <= 1e-6:
            return {"intensity": "Not available", "location": "No focal region detected",
                    "note": "Grad-CAM could not generate a reliable attention map for this image."}

        hot_fraction = float((h >= 0.6 * h.max()).sum()) / h.size

        if hot_fraction < 0.06:
            intensity, intensity_note = "High & focal", "the model's attention is tightly concentrated in a small region"
        elif hot_fraction < 0.18:
            intensity, intensity_note = "Moderate", "the model's attention is moderately concentrated"
        else:
            intensity, intensity_note = "Low & diffuse", "the model's attention is spread across a broad area"

        if not xray_positive:
            return {
                "intensity": "Normal",
                "location": "No significant abnormal focus",
                "note": ("The Grad-CAM visualization shows where the AI model focused while confirming "
                         "a normal chest X-ray. Highlighted regions represent model attention only and "
                         "should not be interpreted as disease.")
            }

        def quadrant_for_point(px, py):
            half = "upper" if py < h.shape[0] / 2 else "lower"
            side = "right" if px < h.shape[1] / 2 else "left"
            return f"{half} zone of the {side} lung field"

        if bbox is not None:
            x1, y1, x2, y2 = bbox
            location = quadrant_for_point((x1 + x2) / 2, (y1 + y2) / 2)
        else:
            py, px = np.unravel_index(np.argmax(h), h.shape)
            location = quadrant_for_point(px, py)

        note = (f"{intensity_note.capitalize()}, concentrated mainly in the {location}. "
                "This is the region that contributed most to the pneumonia prediction. "
                "Highlighted regions show model attention only and must be confirmed by a radiologist.")

        return {"intensity": intensity, "location": location.capitalize(), "note": note}

    except Exception:
        return {"intensity": "Not available", "location": "Not available",
                "note": "Heatmap interpretation could not be generated for this image."}


# =========================
# DATABASE
# =========================
def save_patient(full_name, age, gender):
    conn = get_connection()
    try:
        cursor = conn.cursor()
        try:
            cursor.execute("INSERT INTO patients (full_name, age, gender) VALUES (%s, %s, %s)",
                           (full_name, age, gender))
            conn.commit()
            return cursor.lastrowid
        finally:
            cursor.close()
    finally:
        conn.close()


def save_report(patient_id, report_code, prediction, confidence, image_name, pdf_path, severity_label=None):
    conn = get_connection()
    try:
        cursor = conn.cursor()
        try:
            cursor.execute(
                """INSERT INTO diagnosis_reports
                   (patient_id, report_code, image_name, final_diagnosis, confidence, pdf_path, severity)
                   VALUES (%s, %s, %s, %s, %s, %s, %s)""",
                (patient_id, report_code, image_name, prediction, confidence, pdf_path, severity_label)
            )
            conn.commit()
        finally:
            cursor.close()
    finally:
        conn.close()


# =========================================================================
# EXCEL AUTO-FILL
# =========================================================================
EXCEL_FILE = os.path.join(os.path.dirname(os.path.abspath(__file__)), "patients_data_.xlsx")

EXCEL_FIELDS = ["patient_name", "age", "gender", "wbc", "rbc", "platelet",
                "hb", "temperature", "spo2", "pulse_rate"]

EXCEL_ALIASES = {
    "patient_name": ["patientname", "fullname", "name"],
    "age":          ["age"],
    "gender":       ["gender", "sex"],
    "xray":         ["xrayfile", "xrayimage", "xrayname", "chestxray", "xray", "imagefile",
                     "imagename", "image", "filename", "file"],
    "wbc":          ["wbc", "whitebloodcell"],
    "rbc":          ["rbc", "redbloodcell"],
    "platelet":     ["platelets", "platelet", "plt"],
    "hb":           ["hemoglobin", "haemoglobin", "hgb", "hb"],
    "temperature":  ["temperature", "temp"],
    "spo2":         ["spo2", "oxygen", "o2"],
    "pulse_rate":   ["pulserate", "pulse", "heartrate"],
}


def _excel_key(text):
    return re.sub(r"[^a-z0-9]", "", str(text or "").lower())


def _excel_field_for_header(header):
    k = _excel_key(header)
    if not k:
        return None
    for field, names in EXCEL_ALIASES.items():
        for n in names:
            if k.startswith(n):
                return field
    return None


def _excel_fmt(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def _excel_norm_name(s):
    return " ".join(str(s or "").lower().split())


def _excel_norm_file(s):
    s = os.path.basename(str(s or "").replace("\\", "/")).strip().lower()
    for ext in (".png", ".jpg", ".jpeg"):
        if s.endswith(ext):
            s = s[: -len(ext)]
    return s


def load_excel_rows():
    wb = load_workbook(EXCEL_FILE, data_only=True)
    try:
        ws = wb.active
        all_rows = list(ws.iter_rows(values_only=True))
    finally:
        wb.close()

    if not all_rows:
        return []
    columns = [_excel_field_for_header(h) for h in all_rows[0]]
    if "xray" not in columns:
        raise ValueError("The Excel file needs a column named 'X-ray File'.")

    rows = []
    for raw in all_rows[1:]:
        row = {}
        for field, value in zip(columns, raw):
            if field and field not in row:
                row[field] = _excel_fmt(value)
        if any(row.values()):
            rows.append(row)
    return rows


def save_to_excel(xray_filename, patient_name, age, gender, wbc, rbc, platelet,
                  hb, temperature, spo2, pulse_rate):
    """Save (or update) one patient's row in patients_data_.xlsx."""
    from openpyxl import Workbook

    xray_filename = os.path.basename(str(xray_filename or "")).strip()
    if not xray_filename:
        return

    values_by_field = {
        "patient_name": patient_name, "age": age, "gender": gender,
        "wbc": wbc, "rbc": rbc, "platelet": platelet, "hb": hb,
        "temperature": temperature, "spo2": spo2, "pulse_rate": pulse_rate,
    }

    if os.path.isfile(EXCEL_FILE):
        wb = load_workbook(EXCEL_FILE)
        ws = wb.active
        header_cells = next(ws.iter_rows(min_row=1, max_row=1))
        columns = [_excel_field_for_header(c.value) for c in header_cells]
        if "xray" not in columns:
            return
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = "Patients"
        columns = ["patient_name", "age", "gender", "xray", "wbc", "rbc",
                   "platelet", "hb", "temperature", "spo2", "pulse_rate"]
        ws.append(["Patient Name", "Age", "Gender", "X-ray File",
                   "WBC (cells/uL)", "RBC (mill/uL)", "Platelets (x10^3/uL)",
                   "Hemoglobin (g/dL)", "Temperature (F)", "SpO2 (%)",
                   "Pulse Rate (bpm)"])

    xray_col_index = columns.index("xray") + 1
    target_row = None
    for row in ws.iter_rows(min_row=2):
        cell = row[xray_col_index - 1]
        if _excel_norm_file(cell.value) == _excel_norm_file(xray_filename):
            target_row = row[0].row
            break
    if target_row is None:
        target_row = ws.max_row + 1

    for col_index, field in enumerate(columns, start=1):
        if field == "xray":
            ws.cell(row=target_row, column=col_index, value=xray_filename)
        elif field:
            ws.cell(row=target_row, column=col_index, value=values_by_field.get(field, ""))

    try:
        wb.save(EXCEL_FILE)
    except PermissionError:
        print("Could not save patients_data_.xlsx: please close it in Excel and try again.")


@app.get("/excel-lookup")
def excel_lookup(name: str = "", file: str = ""):
    """Called by autofill.js as soon as an X-ray (and/or name) is entered."""
    fkey = _excel_norm_file(file)
    if not fkey:
        return JSONResponse({"found": False, "message": "Please choose the X-ray image first."})

    if not os.path.isfile(EXCEL_FILE):
        return JSONResponse({"found": False,
                             "message": "patients_data_.xlsx was not found in the Backend folder."})

    try:
        rows = load_excel_rows()
    except Exception as e:
        return JSONResponse({"found": False, "message": f"Could not read patients_data_.xlsx: {e}"})

    by_file = [r for r in rows if _excel_norm_file(r.get("xray")) == fkey]
    if not by_file:
        return JSONResponse({"found": False,
                             "message": f"'{file}' was not found in the X-ray File column of patients_data_.xlsx."})

    nkey = _excel_norm_name(name)
    if nkey:
        matches = [r for r in by_file if _excel_norm_name(r.get("patient_name")) == nkey]
        if not matches:
            excel_name = by_file[0].get("patient_name") or "(no name)"
            return JSONResponse({"found": False,
                                 "message": f"This X-ray is saved in the Excel under the name "
                                            f"'{excel_name}', but you typed '{name.strip()}'. Please check the name."})
    else:
        matches = by_file

    row = matches[0]
    return JSONResponse({"found": True,
                         "message": "Details auto-filled from patients_data_.xlsx.",
                         "data": {f: row.get(f, "") for f in EXCEL_FIELDS}})


# =========================
# STYLES
# =========================
BASE_CSS = """
:root{
  --bg-1:#0f172a; --bg-2:#1e293b;
  --accent:#0ea5e9; --accent-2:#22d3ee;
  --danger:#ef4444; --success:#10b981;
  --card:#ffffff; --text-dark:#0f172a; --text-muted:#64748b;
  --radius:18px;
}
*{box-sizing:border-box;}
body{
  margin:0; min-height:100vh;
  font-family:'Inter','Segoe UI',system-ui,-apple-system,sans-serif;
  background:
    radial-gradient(circle at 15% 20%, rgba(14,165,233,0.18), transparent 45%),
    radial-gradient(circle at 85% 80%, rgba(34,211,238,0.14), transparent 45%),
    linear-gradient(160deg, var(--bg-1), var(--bg-2));
  color:var(--text-dark); padding:36px 18px;
}
.shell{ max-width:1080px; margin:0 auto; }
.brand{ display:flex; align-items:center; justify-content:center; gap:12px; margin-bottom:28px; }
.brand .logo{
  width:46px;height:46px;border-radius:14px;
  background:linear-gradient(135deg,var(--accent),var(--accent-2));
  display:flex;align-items:center;justify-content:center; font-size:22px;
  box-shadow:0 8px 20px rgba(14,165,233,0.45);
}
.brand h1{ color:#f8fafc; font-size:22px; font-weight:700; letter-spacing:-0.01em; margin:0; }
.brand span{ display:block; color:#94a3b8; font-size:12.5px; font-weight:500; margin-top:2px; }

.card{
  background:var(--card); border-radius:var(--radius);
  box-shadow:0 25px 60px -15px rgba(0,0,0,0.4), 0 1px 0 rgba(255,255,255,0.4) inset;
  overflow:hidden;
}
.card-pad{ padding:40px; }
@media (max-width:640px){ .card-pad{ padding:24px; } }

.section-label{
  font-size:14px; font-weight:700; letter-spacing:0.08em; text-transform:uppercase;
  color:var(--accent); margin:0 0 16px 0; display:flex; align-items:center; gap:8px;
}
.section-label::before{ content:""; width:18px;height:2px; background:var(--accent); border-radius:2px; }
.field-group{ margin-bottom:28px; }
.grid{ display:grid; gap:16px; }
.grid-3{ grid-template-columns:repeat(3,1fr); }
.grid-4{ grid-template-columns:repeat(4,1fr); }
@media (max-width:640px){ .grid-3,.grid-4{ grid-template-columns:1fr 1fr; } }

label{ display:block; font-size:13px; font-weight:700; color:#334155; margin-bottom:6px; }
input[type=text], input[type=number], input{
  width:100%; padding:12px 14px; border:1.5px solid #e2e8f0; border-radius:11px;
  font-size:14.5px; background:#f8fafc; transition:border-color .15s, box-shadow .15s; font-family:inherit;
}
input:focus{ outline:none; border-color:var(--accent); box-shadow:0 0 0 4px rgba(14,165,233,0.12); background:#fff; }

.upload-box{
  border:2px dashed #cbd5e1; border-radius:14px; padding:26px; text-align:center;
  background:#f8fafc; transition:border-color .15s, background .15s; cursor:pointer;
}
.upload-box:hover{ border-color:var(--accent); background:#f0f9ff; }
.upload-box input[type=file]{ border:none; background:none; padding:0; text-align:center; }
.upload-icon{ font-size:30px; margin-bottom:6px; }
.upload-hint{ font-size:12.5px; color:var(--text-muted); margin-top:6px; }

.btn{
  display:inline-flex; align-items:center; justify-content:center; gap:8px;
  width:100%; padding:16px; font-size:15.5px; font-weight:700; border:none; border-radius:13px;
  cursor:pointer; color:#fff; background:linear-gradient(135deg,var(--accent),var(--accent-2));
  box-shadow:0 12px 24px -8px rgba(14,165,233,0.55); transition:transform .12s, box-shadow .12s;
}
.btn:hover{ transform:translateY(-1px); box-shadow:0 16px 28px -8px rgba(14,165,233,0.6); }
.btn:active{ transform:translateY(0); }
.btn-outline{ background:#fff; color:var(--accent); border:1.5px solid var(--accent); box-shadow:none; }
.footer-note{ text-align:center; color:#94a3b8; font-size:12px; margin-top:22px; }

/* ---------- Result banner ---------- */
.result-banner{ border-radius:16px; padding:28px; text-align:center; margin-bottom:24px; }
.result-banner.positive{ background:linear-gradient(135deg,#fef2f2,#fee2e2); border:1.5px solid #fecaca; }
.result-banner.negative{ background:linear-gradient(135deg,#f0fdf4,#dcfce7); border:1.5px solid #bbf7d0; }
.result-banner.caution{ background:linear-gradient(135deg,#fffbeb,#fef3c7); border:1.5px solid #fde68a; }
.result-banner h2{ margin:0 0 6px 0; font-size:24px; letter-spacing:-0.01em; }
.result-banner.positive h2{ color:#b91c1c; }
.result-banner.negative h2{ color:#15803d; }
.result-banner.caution h2{ color:#b45309; }
.result-banner .conf{ font-size:14px; color:#475569; font-weight:600; }
.result-banner .sub{ font-size:12.5px; color:#64748b; margin-top:4px; }
.confidence-bar{ width:100%; height:10px; border-radius:6px; background:#e2e8f0; margin-top:14px; overflow:hidden; }
.confidence-fill{ height:100%; border-radius:6px; }
.confidence-fill.positive{ background:linear-gradient(90deg,#f87171,#ef4444); }
.confidence-fill.negative{ background:linear-gradient(90deg,#34d399,#10b981); }
.confidence-fill.caution{ background:linear-gradient(90deg,#fbbf24,#d97706); }

/* ---------- Severity (patient view) ---------- */
.severity-banner{ border-radius:16px; padding:24px 26px; margin-bottom:24px; }
.severity-banner.scap{ background:linear-gradient(135deg,#fef2f2,#fee2e2); border:1.5px solid #fecaca; }
.severity-banner.nscap{ background:linear-gradient(135deg,#fffbeb,#fef3c7); border:1.5px solid #fde68a; }
.severity-banner.unavailable{ background:#f8fafc; border:1.5px dashed #cbd5e1; }
.severity-banner .sev-kicker{ font-size:11.5px; font-weight:700; letter-spacing:0.06em; text-transform:uppercase; color:#64748b; margin-bottom:4px; }
.severity-banner h3{ margin:0 0 8px 0; font-size:21px; letter-spacing:-0.01em; }
.severity-banner h3 small{ font-size:12px; font-weight:600; color:#64748b; margin-left:6px; }
.severity-banner.scap h3{ color:#b91c1c; }
.severity-banner.nscap h3{ color:#b45309; }
.severity-banner.unavailable h3{ color:#64748b; }
.severity-banner .sev-desc{ font-size:14px; color:#334155; line-height:1.6; }

/* ---------- Simple summary card ---------- */
.stage1-card{ background:#f8fafc; border:1px solid #e2e8f0; border-radius:14px; padding:18px 20px; margin-bottom:24px; }
.stage1-card h4{ margin:0 0 10px 0; font-size:13px; letter-spacing:0.04em; text-transform:uppercase; color:var(--accent); }
.stage1-vote-line{ font-size:12.5px; color:#64748b; margin-bottom:12px; }
.stage1-row{
  display:grid; grid-template-columns:200px 160px 1fr; align-items:center; gap:10px;
  padding:9px 0; border-top:1px solid #e2e8f0;
}
.stage1-row:first-of-type{ border-top:none; }
@media (max-width:640px){ .stage1-row{ grid-template-columns:1fr; gap:2px; } }
.stage1-name{ font-size:13px; font-weight:700; color:#334155; }
.stage1-pill{ display:inline-block; padding:3px 10px; border-radius:999px; font-size:11.5px; font-weight:700; width:fit-content; }
.stage1-pill.pos{ background:#fee2e2; color:#b91c1c; }
.stage1-pill.neg{ background:#dcfce7; color:#15803d; }
.stage1-pill.warn{ background:#fef3c7; color:#b45309; }
.stage1-pill.na{ background:#e2e8f0; color:#64748b; }
.stage1-detail{ font-size:12px; color:#64748b; }

.info-note{ border-radius:14px; padding:18px 20px; margin-bottom:24px; border:1.5px solid #fde68a; background:linear-gradient(135deg,#fffbeb,#fef3c7); }
.info-note h4{ margin:0 0 6px 0; font-size:14px; color:#b45309; }
.info-note p{ margin:0; font-size:13px; color:#334155; line-height:1.6; }

.image-panel{ background:#0f172a; border-radius:16px; padding:18px; text-align:center; margin-bottom:24px; }
.image-panel img{ max-width:100%; border-radius:10px; box-shadow:0 8px 30px rgba(0,0,0,0.4); }
.image-panel .cap{ color:#94a3b8; font-size:12px; margin-top:10px; letter-spacing:0.03em; }

.info-grid{ display:grid; grid-template-columns:repeat(3,1fr); gap:14px; margin-bottom:24px; }
@media (max-width:640px){ .info-grid{ grid-template-columns:1fr 1fr; } }
.info-tile{ background:#f8fafc; border:1px solid #e2e8f0; border-radius:12px; padding:14px 16px; }
.info-tile .k{ font-size:11px; color:var(--text-muted); font-weight:700; text-transform:uppercase; letter-spacing:0.05em; }
.info-tile .v{ font-size:16px; font-weight:700; color:var(--text-dark); margin-top:3px; }

.lab-card{ background:#fff; border:1px solid #e2e8f0; border-radius:14px; overflow:hidden; margin-bottom:24px; overflow-x:auto; }
.lab-table{ width:100%; border-collapse:collapse; font-size:13.5px; }
.lab-table thead th{
  background:linear-gradient(135deg,var(--accent),var(--accent-2)); color:#fff; text-align:left;
  padding:12px 14px; font-size:11.5px; letter-spacing:0.04em; text-transform:uppercase;
}
.lab-table tbody td{ padding:11px 14px; border-bottom:1px solid #f1f5f9; color:#0f172a; }
.lab-table tbody tr:last-child td{ border-bottom:none; }
.lab-table tbody tr:nth-child(even){ background:#f8fafc; }
.lab-table .range-col{ color:var(--text-muted); }
.flag-pill{ display:inline-block; padding:3px 10px; border-radius:999px; font-size:11.5px; font-weight:700; }
.flag-pill.Normal{ background:#dcfce7; color:#15803d; }
.flag-pill.Low{ background:#fef3c7; color:#b45309; }
.flag-pill.High{ background:#fee2e2; color:#b91c1c; }
.flag-pill.NA{ background:#e2e8f0; color:#64748b; }

/* ---------- Abnormal values note + Recommendation ---------- */
.note-heading{ font-size:14px; font-weight:800; letter-spacing:0.02em; text-transform:uppercase; color:#0f172a;
  border-bottom:1px solid #e2e8f0; padding-bottom:6px; margin:0 0 12px 0; }
.note-box{ border-radius:6px; padding:16px 20px; font-size:14px; line-height:1.6; color:#334155;
  margin-bottom:24px; border:1px solid; border-left-width:6px; }
.note-box.amber{ background:#fffbeb; border-color:#fde68a; border-left-color:#d97706; }
.note-box.red{ background:#fef2f2; border-color:#fecaca; border-left-color:#dc2626; }
.note-box.green{ background:#f0fdf4; border-color:#bbf7d0; border-left-color:#16a34a; }

/* ---------- Doctor-only collapsible ---------- */
.doctor-details{ border:1px solid #e2e8f0; border-radius:14px; background:#fff; margin-bottom:24px; overflow:hidden; }
.doctor-details > summary{
  cursor:pointer; list-style:none; padding:16px 20px; font-size:14px; font-weight:700; color:#334155;
  display:flex; align-items:center; justify-content:space-between; background:#f8fafc;
}
.doctor-details > summary::-webkit-details-marker{ display:none; }
.doctor-details > summary::after{ content:"\\25BE"; color:var(--accent); font-size:16px; transition:transform .15s; }
.doctor-details[open] > summary::after{ transform:rotate(180deg); }
.doctor-details > summary span{ font-weight:500; font-size:12px; color:#94a3b8; margin-left:8px; }
.doctor-body{ padding:20px; }
.doctor-body .stage1-card, .doctor-body .interp-card, .doctor-body .glossary-card{ margin-bottom:18px; }
.doctor-body h5{ margin:18px 0 8px 0; font-size:12px; letter-spacing:0.05em; text-transform:uppercase; color:var(--accent); }
.doctor-body h5:first-child{ margin-top:0; }

.tech-row{ display:grid; grid-template-columns:130px 110px 1fr; gap:10px; padding:8px 0; border-top:1px solid #e2e8f0; align-items:center; }
.tech-row:first-of-type{ border-top:none; }
@media (max-width:560px){ .tech-row{ grid-template-columns:1fr; gap:2px; } }

.sev-breakdown-row{ display:grid; grid-template-columns:230px 1fr 170px; align-items:center; gap:10px; margin-bottom:6px; }
@media (max-width:640px){ .sev-breakdown-row{ grid-template-columns:1fr; gap:3px; } }
.sev-breakdown-name{ font-size:12px; font-weight:700; color:#334155; }
.sev-breakdown-bar{ height:8px; border-radius:4px; background:rgba(0,0,0,0.08); overflow:hidden; display:block; }
.sev-breakdown-bar span{ display:block; height:100%; border-radius:4px; background:#0ea5e9; }
.sev-breakdown-val{ font-size:11px; color:#64748b; text-align:right; }
@media (max-width:640px){ .sev-breakdown-val{ text-align:left; } }

.interp-card{ background:#f8fafc; border:1px solid #e2e8f0; border-radius:14px; padding:16px 18px; }
.interp-grid{ display:grid; grid-template-columns:1fr 1fr; gap:14px; margin-bottom:12px; }
@media (max-width:520px){ .interp-grid{ grid-template-columns:1fr; } }
.interp-tile .k{ font-size:11px; color:var(--text-muted); font-weight:700; text-transform:uppercase; letter-spacing:0.05em; }
.interp-tile .v{ font-size:15px; font-weight:700; color:var(--text-dark); margin-top:4px; }
.interp-note{ font-size:13px; color:#334155; line-height:1.6; padding-top:10px; border-top:1px solid #e2e8f0; }

.glossary-card{ background:#f8fafc; border:1px solid #e2e8f0; border-radius:14px; padding:16px 18px; }
.glossary-card dl{ margin:0; }
.glossary-card dt{ font-size:13.5px; font-weight:700; color:#0f172a; margin-top:10px; }
.glossary-card dt:first-child{ margin-top:0; }
.glossary-card dd{ margin:2px 0 0 0; font-size:13px; color:#334155; line-height:1.6; }

.actions{ display:flex; gap:12px; justify-content:center; align-items:center; margin-top:10px; }
.actions .btn{ width:auto; min-width:160px; text-decoration:none; text-align:center; }
.actions .btn-outline{ background:#fff; color:var(--accent); border:1.5px solid var(--accent); box-shadow:none; }
.actions .btn-outline:hover{ background:var(--accent); color:#fff; transform:translateY(-1px); box-shadow:0 8px 18px rgba(14,165,233,0.25); }
@media (max-width:520px){ .actions{ flex-direction:column; } .actions .btn{ width:100%; } }

.disclaimer{ margin-top:26px; font-size:12px; color:#94a3b8; text-align:center; line-height:1.6; }
"""


# =========================
# HOME PAGE
# =========================
@app.get("/", response_class=HTMLResponse)
def home(patient_name: str = "", age: str = "", gender: str = ""):
    patient_name_val = html.escape(patient_name)
    age_val = html.escape(age)
    gender_val = html.escape(gender)
    return f"""
<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>Pneumonia Detection System &middot; X-Ray Screening</title>
<style>{BASE_CSS}</style>
</head>
<body>

<div class="shell">

  <div class="brand">
    <div class="logo">🫁</div>
    <h1>Pneumonia Detection System
      <span>Chest X-Ray Pneumonia Screening &amp; Report Generator</span>
    </h1>
  </div>

  <div class="card">
    <div class="card-pad">

      <form action="/predict" method="post" enctype="multipart/form-data">

        <div class="field-group">
          <p class="section-label">Patient Information</p>
          <div class="grid grid-3">
            <div>
              <label>Patient Name</label>
              <input type="text" name="patient_name" placeholder="e.g. John Carter" value="{patient_name_val}" required>
            </div>
            <div>
              <label>Age</label>
              <input type="number" name="age" placeholder="e.g. 34" value="{age_val}" required>
            </div>
            <div>
              <label>Gender</label>
              <input type="text" name="gender" placeholder="Male / Female / Other" value="{gender_val}" required>
            </div>
          </div>
        </div>

        <div class="field-group">
          <p class="section-label">Chest X-Ray Image</p>
          <div class="upload-box">
            <div class="upload-icon">📤</div>
            <input type="file" name="file" accept="image/*" required>
            <div class="upload-hint">JPG or PNG &middot; the image will be analyzed by the AI model</div>
          </div>
        </div>

        <div class="field-group">
          <p class="section-label">Lab Values</p>
          <div class="grid grid-4">
            <div>
              <label>WBC (cells/µL)</label>
              <input type="text" name="wbc" placeholder="e.g. 8200">
            </div>
            <div>
              <label>RBC (mill/µL)</label>
              <input type="text" name="rbc" placeholder="e.g. 4.8">
            </div>
            <div>
              <label>Platelets (×10³/µL)</label>
              <input type="text" name="platelet" placeholder="e.g. 250">
            </div>
            <div>
              <label>Hemoglobin (g/dL)</label>
              <input type="text" name="hb" placeholder="e.g. 13.5">
            </div>
          </div>
        </div>

        <div class="field-group">
          <p class="section-label">Vitals</p>
          <div class="grid grid-3">
            <div>
              <label>Temperature (°F)</label>
              <input type="text" name="temperature" placeholder="e.g. 98.6">
            </div>
            <div>
              <label>SpO2 (%)</label>
              <input type="text" name="spo2" placeholder="e.g. 97">
            </div>
            <div>
              <label>Pulse Rate (bpm)</label>
              <input type="text" name="pulse_rate" placeholder="e.g. 78">
            </div>
          </div>
        </div>

        <button class="btn" type="submit">Analyze &amp; Generate Report</button>

      </form>

    </div>
  </div>

  <div style="max-width:1080px;margin:18px auto 0 auto;">
    <a href="/history" class="btn btn-outline" style="display:flex;">Stored Records</a>
  </div>
  <p class="footer-note">For clinical decision support only &middot; not a substitute for radiologist review</p>

</div>

<script src="/assets/autofill.js"></script>
</body>
</html>
"""


# =========================
# HTML BUILDERS
# =========================
GLOSSARY_INTRO_TEXT = (
    "Severity staging follows the ATS/IDSA 2007 criteria for severe community-acquired "
    "pneumonia: a case is SCAP when 3 or more minor criteria are met. Only the criteria "
    "that can be evaluated from the entered values are scored (WBC, Platelets, Temperature, "
    "SpO2 as a proxy for hypoxemia). Grad-CAM is used only as a visual explanation."
)


def build_glossary_html():
    nscap = SEVERITY_INFO["NSCAP"]
    scap = SEVERITY_INFO["SCAP"]
    return f"""
      <div class="glossary-card">
        <p style="font-size:12.5px;color:#64748b;margin:0 0 12px 0;line-height:1.6;">{GLOSSARY_INTRO_TEXT}</p>
        <dl>
          <dt>NSCAP &mdash; {nscap['full_name']}</dt>
          <dd>{nscap['description']}</dd>
          <dt>SCAP &mdash; {scap['full_name']}</dt>
          <dd>{scap['description']}</dd>
        </dl>
      </div>"""


def build_stage1_patient_html(stage1_result):
    rows = ""
    oxygen_grade = stage1_result.get("oxygen_grade")
    blood_abnormal = stage1_result.get("blood_abnormal")
    urgent_override = stage1_result.get("urgent_override", False)
    for c in stage1_result["components"]:
        name, text, pill_class, note = friendly_signal(c, oxygen_grade, blood_abnormal, urgent_override)
        rows += f"""
        <div class="stage1-row">
          <span class="stage1-name">{html.escape(name)}</span>
          <span class="stage1-pill {pill_class}">{html.escape(text)}</span>
          <span class="stage1-detail">{html.escape(note)}</span>
        </div>"""

    disagree_note = ("<div class='stage1-vote-line' style='margin:12px 0 0 0;'><strong>Some checks "
                     "disagree with each other, so a doctor's review is strongly advised.</strong></div>"
                     if stage1_result["disagree"] else "")
    return f"""
      <div class="stage1-card">
        <h4>Summary of Checks</h4>
        {rows}
        {disagree_note}
      </div>"""


def build_stage1_technical_html(stage1_result):
    rows = ""
    blood_abnormal = stage1_result.get("blood_abnormal")
    for c in stage1_result["components"]:
        pill_class, pill_text, detail = technical_signal(c, blood_abnormal, stage1_result.get("oxygen_grade"))
        rows += f"""
        <div class="tech-row">
          <span class="stage1-name">{html.escape(c['name'])}</span>
          <span class="stage1-pill {pill_class}">{html.escape(pill_text)}</span>
          <span class="stage1-detail">{html.escape(detail)}</span>
        </div>"""

    og = stage1_result.get("oxygen_grade")
    og_html = ""
    if og and og.get("available"):
        og_html = (f"<div class='stage1-vote-line' style='margin:12px 0 0 0;'>"
                   f"Oxygen &amp; pulse grading ({html.escape(og['situation_label'])} SpO2 thresholds): "
                   f"{html.escape(og['breakdown'])} &rarr; <strong>{html.escape(og['label'])}</strong> "
                   f"(worst value decides)</div>")
    if stage1_result.get("urgent_override"):
        og_html += ("<div class='stage1-vote-line' style='margin:12px 0 0 0;color:#b45309;'>"
                    "<strong>Safety rule applied:</strong> no pneumonia pattern was found, "
                    "but the Oxygen &amp; pulse grade is Abnormal, so the headline "
                    "shows <strong>urgent medical review advised</strong> instead of an unqualified "
                    "&ldquo;no signs&rdquo; message.</div>")

    reasons_html = "".join(f"<li>{html.escape(r)}</li>" for r in stage1_result.get("reasons", []))
    reasons_block = (f"<ul style='margin:6px 0 10px 18px;padding:0;font-size:12.5px;color:#334155;'>"
                     f"{reasons_html}</ul>") if reasons_html else ""

    return f"""
      <div class="stage1-card">
        <div class="stage1-vote-line">
          Decision rule: the X-ray model is the primary signal and a positive X-ray is never overridden
          by the other models. Combined pneumonia likelihood (weighted, X-ray highest):
          <strong>{stage1_result['final_prob'] * 100:.1f}%</strong> &rarr;
          <strong>{html.escape(stage1_result['display_label'])}</strong>
        </div>
        {reasons_block}
        {rows}
        {og_html}
      </div>"""


def build_severity_patient_html(severity_info):
    if not (severity_info and severity_info["available"]):
        reason = html.escape(severity_info["reason"]) if severity_info else "Severity could not be assessed."
        return f"""
      <div class="severity-banner unavailable">
        <div class="sev-kicker">Severity check</div>
        <h3>Not assessed</h3>
        <div class="sev-desc">{reason}</div>
      </div>"""

    p = severity_info["patient"]
    sev_class = "scap" if severity_info["level"] == "severe" else "nscap"
    provisional_note = ""
    if severity_info["provisional"] and severity_info["level"] == "mild":
        provisional_note = (" Some checks could not be done because the values were not entered, "
                            "so this result is only a first estimate.")
    return f"""
      <div class="severity-banner {sev_class}">
        <div class="sev-kicker">Severity check</div>
        <h3>{p['title']} <small>({p['tag']})</small></h3>
        <div class="sev-desc">{p['text']}{provisional_note}</div>
      </div>"""


def build_severity_technical_html(severity_info):
    if not (severity_info and severity_info["available"]):
        return ""
    rows = "".join(
        f"""<div class="sev-breakdown-row">
            <span class="sev-breakdown-name">{html.escape(b['name'])}</span>
            <span class="sev-breakdown-bar"><span style="width:{100 if b['met'] else 0}%;"></span></span>
            <span class="sev-breakdown-val">{html.escape(b['status'])} &middot; {html.escape(b['detail'])}</span>
        </div>"""
        for b in severity_info["criteria"]
    )
    prov = " (provisional)" if severity_info["provisional"] else ""
    override = ("<p style='font-size:12.5px;color:#b91c1c;margin:8px 0 0 0;'><strong>Safety rule applied:</strong> "
                "an Abnormal (clearly low) SpO2 for this patient's situation triggers the urgent message "
                "regardless of the criteria count.</p>"
                if severity_info.get("urgent_override") and severity_info["label"] != "SCAP" else "")
    return f"""
      <h5>Severity staging &middot; {severity_info['label']}{prov} &middot; {html.escape(severity_info['full_name'])}</h5>
      <p style="font-size:13px;color:#334155;line-height:1.6;margin:0 0 10px 0;">{severity_info['description']}</p>
      <p style="font-size:12.5px;color:#64748b;margin:0 0 10px 0;">
        {severity_info['met']} of {severity_info['evaluated']} evaluable ATS/IDSA 2007 minor criteria met (SCAP needs 3)
      </p>
      {rows}
      {override}"""


SHORT_LAB_NAMES = {"Temperature": "Temp", "Pulse Rate": "Pulse", "Hemoglobin": "Hb"}


def _flag_phrase(r, lab_rows, spo2_grade):
    """
    One consistent wording per value. Temperature, pulse and SpO2 use the same Normal /
    Borderline / Abnormal grade as the 'Oxygen & pulse' row, so the report never calls a value
    'high' in one place and 'borderline' in another. Returns None when the value is not listed.
    """
    label = r["label"]
    short = SHORT_LAB_NAMES.get(label, label)
    direction = r["flag"].lower()            # "high" / "low"
    val = _lab_value(lab_rows, label)

    grade = None
    if label == "SpO2":
        grade = spo2_grade
    elif label == "Pulse Rate" and val is not None:
        grade = grade_pulse(val)
    elif label == "Temperature" and val is not None:
        grade = grade_temperature_f(val)

    if grade == "Normal":
        return None
    if grade == "Borderline":
        return f"{short} slightly {direction}"
    return f"{short} {direction}"


def build_value_notes(is_positive, lab_rows, stage1_result, severity_info, age=None):
    """Condition/situation-based text for 'ABNORMAL VALUES' and 'RECOMMENDATION' (web + PDF)."""
    flag = {r["label"]: r["flag"] for r in lab_rows}
    verdict = (stage1_result or {}).get("verdict")

    # Situation-based SpO2 grade (Normal / Borderline / Abnormal)
    og = (stage1_result or {}).get("oxygen_grade") or {}
    spo2_grade = og.get("spo2_grade")

    abnormal = []
    for r in lab_rows:
        if r["flag"] not in ("High", "Low"):
            continue
        phrase = _flag_phrase(r, lab_rows, spo2_grade)
        if phrase:
            abnormal.append(phrase)

    spo2 = _lab_value(lab_rows, "SpO2")
    temp = _lab_value(lab_rows, "Temperature")
    if spo2_grade is not None:
        low_o2 = spo2_grade == "Abnormal"
        borderline_o2 = spo2_grade == "Borderline"
    else:
        low_o2 = spo2 is not None and spo2 < SPO2_HYPOXEMIA_CUTOFF
        borderline_o2 = spo2 is not None and SPO2_HYPOXEMIA_CUTOFF <= spo2 < 95
    high_fever = temp is not None and temp >= 102
    fast_pulse = flag.get("Pulse Rate") == "High"
    severe = bool(severity_info and severity_info.get("available") and severity_info.get("level") == "severe")
    disagree = bool(stage1_result and stage1_result.get("disagree"))
    n_missing = sum(1 for r in lab_rows if r["flag"] == "N/A")
    try:
        age_val = float(age)
    except (TypeError, ValueError):
        age_val = None
    high_risk_age = age_val is not None and (age_val >= 65 or age_val < 5)

    # ---------- PLEASE NOTE: ABNORMAL VALUES ----------
    abnormal_text = None
    if abnormal:
        listed = ", ".join(abnormal)
        if verdict == LIKELY or (is_positive and verdict is None):
            text = (f"The screening suggests pneumonia, and some entered values are outside the usual "
                    f"range ({listed}).")
        elif verdict == POSSIBLE and (stage1_result or {}).get("driver") == "blood":
            text = (f"The chest X-ray looks normal, but the blood results need review, and some entered "
                    f"values are outside the usual range ({listed}).")
        elif verdict == POSSIBLE:
            text = (f"The screening shows findings that may indicate pneumonia, and some entered values are "
                    f"outside the usual range ({listed}).")
        elif verdict == INCONCLUSIVE:
            text = (f"The checks could not be combined into a reliable result, and some entered values are "
                    f"outside the usual range ({listed}).")
        else:
            text = (f"The screening did not find evidence of pneumonia, but some entered values are outside "
                    f"the usual range ({listed}). These can have causes other than pneumonia, and a normal "
                    f"screening result does not rule out illness.")

        meanings = []
        if flag.get("WBC") == "High":
            meanings.append("Raised WBC can point to infection or inflammation.")
        if flag.get("WBC") == "Low":
            meanings.append("A low WBC can reflect a viral illness or weakened immunity.")
        if high_fever or flag.get("Temperature") == "High":
            meanings.append("A raised temperature indicates fever" +
                            (" and a fast pulse often goes with it." if fast_pulse else "."))
        elif fast_pulse:
            meanings.append("A fast pulse can be caused by fever, low oxygen, pain or anxiety.")
        if flag.get("Hemoglobin") == "Low" or flag.get("RBC") == "Low":
            meanings.append("Low Hb/RBC suggests anaemia, which can add to tiredness and breathlessness.")
        if flag.get("Platelets") == "Low":
            meanings.append("Low platelets can increase bleeding risk and should be reviewed.")
        elif flag.get("Platelets") == "High":
            meanings.append("High platelets can occur with infection or inflammation.")
        if all(flag.get(k) == "Low" for k in ("WBC", "RBC", "Platelets", "Hemoglobin")):
            meanings.append("All blood counts are low together, which is not a typical pneumonia pattern "
                            "and needs urgent evaluation for other causes.")
        if low_o2:
            meanings.append("A low oxygen level needs prompt medical attention.")
        elif borderline_o2:
            meanings.append("The oxygen level is slightly below normal and should be rechecked.")

        text += " " + " ".join(meanings) if meanings else ""
        text += " A medical check-up is advised, especially if symptoms are present."
        abnormal_text = text

    # ---------- RECOMMENDATION ----------
    urgent = low_o2 or severe
    if is_positive:
        if verdict == POSSIBLE and (stage1_result or {}).get("driver") == "blood":
            rec = ("The chest X-ray looks normal, but the blood results need a doctor's review to find "
                   "the cause.")
        elif verdict == POSSIBLE:
            rec = ("Some findings may indicate pneumonia. Please consult a doctor to confirm or exclude it "
                   "and decide on any treatment.")
        else:
            rec = ("The screening suggests pneumonia. Please consult a doctor to confirm the result "
                   "and decide on treatment.")
        if urgent:
            rec += " Please go to a hospital or emergency care today and do not wait for symptoms to worsen."
        else:
            rec += " Please see a doctor within 24 hours."
    elif verdict == INCONCLUSIVE:
        rec = ("The checks could not be combined into a reliable result, so pneumonia can be neither "
               "confirmed nor ruled out. A doctor's review is required.")
        if urgent:
            rec += " Because the oxygen level is low, please get medical attention promptly."
        rec += " Medical attention should be sought if fever, cough or breathlessness persists or worsens."
    else:
        rec = ("No evidence of pneumonia was identified on the available investigations. Clinical "
               "correlation is advised.")
        if urgent:
            rec += (" Because the oxygen level is low, please get medical attention promptly even though "
                    "pneumonia was not detected.")
        elif abnormal:
            rec += (" Some of the entered values are outside the usual range, so a medical check-up is "
                    "advised.")
        rec += " Medical attention should be sought if fever, cough or breathlessness persists or worsens."

    if high_risk_age:
        rec += " Because of the patient's age, please seek care early rather than waiting."
    if borderline_o2 and not low_o2:
        rec += " Recheck the oxygen level and seek care if it drops further."
    if high_fever:
        rec += " A persistent high fever needs medical review."
    if disagree:
        rec += " Individual investigations gave differing results, so clinical correlation is particularly important."
    if n_missing >= 3:
        rec += " Several values were not entered, so this assessment is limited and should be repeated with complete values."
    if urgent or is_positive:
        rec += (" Seek emergency care immediately if breathing becomes difficult, the lips turn bluish, "
                "or there is chest pain or confusion.")

    if urgent:
        tone = "red"
    elif is_positive or abnormal or disagree or high_risk_age or verdict == INCONCLUSIVE:
        tone = "amber"
    else:
        tone = "green"

    return {"abnormal_text": abnormal_text, "recommendation": rec, "tone": tone}


def build_value_notes_html(notes):
    out = ""
    if notes["abnormal_text"]:
        out += f"""
      <div class="note-heading">Please note: abnormal values</div>
      <div class="note-box amber">{notes['abnormal_text']}</div>"""
    out += f"""
      <div class="note-heading">Recommendation</div>
      <div class="note-box {notes['tone']}">{notes['recommendation']}</div>"""
    return out


# =========================
# PREDICT
# =========================
@app.post("/predict", response_class=HTMLResponse)
async def predict(
    patient_name: str = Form(...),
    age: str = Form(...),
    gender: str = Form(...),
    wbc: str = Form(""),
    rbc: str = Form(""),
    platelet: str = Form(""),
    hb: str = Form(""),
    temperature: str = Form(""),
    spo2: str = Form(""),
    pulse_rate: str = Form(""),
    file: UploadFile = File(...)
):
    try:
        if model is None:
            return HTMLResponse("<h2>Model not loaded</h2>")

        cleanup_old_files("static")
        cleanup_old_files("reports")
        cleanup_stale_pdf_store_entries()

        patient_name_safe = html.escape(patient_name)
        age_safe = html.escape(age)
        gender_safe = html.escape(gender)

        contents = await file.read()

        try:
            save_to_excel(
                xray_filename=file.filename,
                patient_name=patient_name, age=age, gender=gender,
                wbc=wbc, rbc=rbc, platelet=platelet, hb=hb,
                temperature=temperature, spo2=spo2, pulse_rate=pulse_rate,
            )
        except Exception as excel_err:
            print("Could not save to patients_data_.xlsx:", excel_err)

        img, img_array = preprocess(contents)

        pred = model.predict(img_array)
        raw_prob = float(pred.squeeze())
        pneumonia_prob = raw_prob if POSITIVE_CLASS_IS_PNEUMONIA_WHEN_PROB_HIGH else (1 - raw_prob)
        print(f"[predict] raw_model_output={raw_prob:.4f}  pneumonia_prob={pneumonia_prob:.4f}  "
              f"polarity_flag={POSITIVE_CLASS_IS_PNEUMONIA_WHEN_PROB_HIGH}")

        vitals = {
            "wbc": wbc, "rbc": rbc, "platelet": platelet, "hb": hb,
            "temperature": temperature, "spo2": spo2, "pulse_rate": pulse_rate
        }
        lab_rows = evaluate_all_vitals(vitals)

        def cleaned(label_name):
            r = next(r for r in lab_rows if r["label"] == label_name)
            return "" if r["flag"] == "N/A" else r["value"].rstrip(" *")

        blood_result = predict_blood_ann(age, cleaned("Hemoglobin"), cleaned("Platelets"),
                                         cleaned("WBC"), cleaned("RBC"), gender)
        oxygen_result = predict_oxygen_ann(cleaned("SpO2"), cleaned("Pulse Rate"), cleaned("Temperature"))

        # ---- Corrected fusion: X-ray is primary; a positive X-ray can never become "no signs" ----
        stage1_result = combine_stage1(pneumonia_prob, blood_result, oxygen_result)

        is_positive = stage1_result["is_positive"]          # LIKELY or POSSIBLE
        xray_positive = stage1_result["xray_positive"]      # the image model's own call
        verdict = stage1_result["verdict"]
        label = stage1_result["label"]                      # technical label (stored in DB)
        confidence = stage1_result["confidence"]            # the ONE score: pneumonia likelihood %
        screening_score = confidence

        # 3-label grade (Normal / Borderline / Abnormal) for "Oxygen & pulse".
        spo2_situation = pick_spo2_situation(age, is_positive)
        oxygen_grade = grade_oxygen_group(lab_rows, spo2_situation)
        stage1_result["oxygen_grade"] = oxygen_grade

        # Safety rule: negative result + Abnormal oxygen/pulse/temperature -> urgent review wording.
        urgent_override = needs_urgent_review(is_positive, oxygen_grade) and verdict == NEGATIVE
        stage1_result["urgent_override"] = urgent_override
        headline = headline_for(stage1_result, urgent_override)
        result_class = result_class_for(stage1_result, urgent_override)

        # Blood values outside the reference range (rule-based).
        blood_abnormal = blood_values_out_of_range(lab_rows)
        stage1_result["blood_abnormal"] = blood_abnormal
        blood_component = next((c for c in stage1_result["components"] if c["name"] == "Blood ANN"), None)
        if (blood_abnormal and blood_component and blood_component["available"]
                and not blood_component["is_positive"]):
            stage1_result["disagree"] = True

        stage1_patient_html = build_stage1_patient_html(stage1_result)
        stage1_tech_html = build_stage1_technical_html(stage1_result)

        # ---- Grad-CAM: driven by the X-RAY model's own prediction, not the fused result ----
        last_conv, layer_owner = get_last_conv(model)
        heatmap = gradcam(model, img_array, last_conv, layer_owner)
        if xray_positive and float(np.max(heatmap)) <= 1e-6:
            print("[QA] X-ray model is positive but Grad-CAM is empty - check the Grad-CAM layer/pipeline.")

        bbox = get_bbox_from_heatmap(heatmap) if xray_positive else None
        region_localized = bbox is not None
        heatmap_info = interpret_heatmap(heatmap, xray_positive, bbox)

        if xray_positive:
            comparison_img = make_comparison_image(img, heatmap, bbox)
        else:
            img_r = cv2.resize(img, (IMG_SIZE, IMG_SIZE))
            comparison_img = _add_caption(img_r, "Original Image")

        img_name = f"{uuid.uuid4().hex}.jpg"
        cv2.imwrite(f"static/{img_name}", comparison_img)

        report_id = uuid.uuid4().hex[:10].upper()
        timestamp = datetime.now().strftime("%d %b %Y, %I:%M %p")

        spo2_abnormal = (oxygen_grade["spo2_grade"] == "Abnormal") if oxygen_grade.get("spo2_grade") else None
        severity_info = predict_severity(lab_rows, spo2_abnormal) if (is_positive and xray_positive) else None

        lab_rows_html = "".join(
            f"""<tr>
                <td>{r['label']}</td>
                <td><strong>{r['value']}</strong></td>
                <td class="range-col">{r['unit']}</td>
                <td class="range-col">{r['range']}</td>
                <td><span class="flag-pill {r['flag'].replace('/', '')}">{r['flag']}</span></td>
            </tr>"""
            for r in lab_rows
        )
        unit_warnings = [r for r in lab_rows if r.get("unit_warning")]
        unit_warning_html = ""
        if unit_warnings:
            items = "".join(f"<li>{r['label']}: value {r['unit_warning']}</li>" for r in unit_warnings)
            unit_warning_html = f"""
            <div style="font-size:12px;color:#b45309;margin-top:-12px;margin-bottom:24px;">
              ⚠ Some entered values looked like they were in a different unit and were
              auto-converted before checking:
              <ul style="margin:4px 0 0 18px;padding:0;">{items}</ul>
              Please double-check the original entry.
            </div>"""

        severity_patient_html = build_severity_patient_html(severity_info) if severity_info is not None else ""
        severity_tech_html = build_severity_technical_html(severity_info) if severity_info is not None else ""
        value_notes_html = build_value_notes_html(
            build_value_notes(is_positive, lab_rows, stage1_result, severity_info, age))

        # CAP / SCAP legend is only relevant when pneumonia is suspected
        glossary_block = ""
        if severity_info is not None:
            glossary_block = f"""
      <div style="margin-bottom:24px;">
        <h5 style="margin:0 0 8px 0;font-size:12px;letter-spacing:0.05em;text-transform:uppercase;color:#0ea5e9;">About severity staging (NSCAP / SCAP)</h5>
        {build_glossary_html()}
      </div>"""

        image_caption = (
            "Yellow box marks the area of concern \u2022 the red/yellow heatmap shows where the AI focused."
            if xray_positive
            else "Original chest X-ray (heatmap hidden because the X-ray model found no abnormal region)."
        )

        pdf_path = build_pdf_report(
            report_id=report_id, timestamp=timestamp,
            patient_name=patient_name, age=age, gender=gender,
            label=label, headline=headline, screening_score=screening_score,
            confidence=confidence, is_positive=is_positive,
            image_path=f"static/{img_name}", image_shape=comparison_img.shape,
            vitals=vitals, lab_rows=lab_rows, heatmap_info=heatmap_info,
            region_localized=region_localized, severity_info=severity_info,
            stage1_result=stage1_result, urgent_override=urgent_override,
            result_class=result_class, xray_positive=xray_positive
        )
        PDF_STORE[report_id] = pdf_path

        try:
            age_for_db = int(age)
        except (TypeError, ValueError):
            age_for_db = None

        try:
            patient_id = save_patient(patient_name, age_for_db, gender)
            save_report(
                patient_id=patient_id, report_code=report_id, prediction=label,
                confidence=round(confidence, 2), image_name=img_name, pdf_path=pdf_path,
                severity_label=(severity_info["label"] if severity_info and severity_info["available"] else None)
            )
        except Exception as db_err:
            print("DB save error:", db_err)

        return f"""
<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>Report Result &middot; Pneumonia Detection System</title>
<style>{BASE_CSS}</style>
</head>
<body>

<div class="shell">

  <div class="brand">
    <div class="logo">🫁</div>
    <h1>Pneumonia Detection System
      <span>Screening Report &middot; #{report_id}</span>
    </h1>
  </div>

  <div class="card">
    <div class="card-pad">

      <div class="info-grid">
        <div class="info-tile"><div class="k">Patient</div><div class="v">{patient_name_safe}</div></div>
        <div class="info-tile"><div class="k">Age / Gender</div><div class="v">{age_safe} / {gender_safe}</div></div>
        <div class="info-tile"><div class="k">Scan Date</div><div class="v">{timestamp}</div></div>
      </div>

      <div class="result-banner {result_class}">
        <h2>{headline}</h2>
        <div class="conf">Pneumonia likelihood: {screening_score:.1f}%</div>
        <div class="confidence-bar">
          <div class="confidence-fill {result_class}" style="width:{screening_score:.1f}%;"></div>
        </div>
        <div class="sub">Combined likelihood from the available checks (X-ray weighted highest). This is a screening result, not a diagnosis.</div>
      </div>

      {severity_patient_html}
      {value_notes_html}
      {stage1_patient_html}

      <div class="lab-card">
        <table class="lab-table">
          <thead>
            <tr><th>Test</th><th>Result</th><th>Unit</th><th>Normal Range</th><th>Status</th></tr>
          </thead>
          <tbody>{lab_rows_html}</tbody>
        </table>
      </div>
      {unit_warning_html}

      <div class="image-panel">
        <img src="/static/{img_name}" alt="Chest X-ray with highlighted area">
        <div class="cap">{image_caption}</div>
      </div>

      {glossary_block}

      <details class="doctor-details">
        <summary>Clinical details <span>(for doctors)</span></summary>
        <div class="doctor-body">
          <h5>Model signals &amp; probabilities</h5>
          {stage1_tech_html}
          {severity_tech_html}
          <h5>Grad-CAM interpretation</h5>
          <div class="interp-card">
            <div class="interp-grid">
              <div class="interp-tile"><div class="k">Attention Intensity</div><div class="v">{heatmap_info['intensity']}</div></div>
              <div class="interp-tile"><div class="k">Region of Focus</div><div class="v">{heatmap_info['location']}</div></div>
            </div>
            <div class="interp-note">{heatmap_info['note']}</div>
          </div>
        </div>
      </details>

      <div class="actions">
        <a href="/download?report_id={report_id}" class="btn">⬇ Download PDF Report</a>
        <a href="#" onclick="window.print(); return false;" class="btn btn-outline">Print</a>
        <a href="{STREAMLIT_URL}" class="btn btn-outline">Go to Home</a>
      </div>

      <p class="disclaimer">{SCREENING_DISCLAIMER}</p>

    </div>
  </div>

</div>

</body>
</html>
"""

    except Exception:
        traceback.print_exc()
        return HTMLResponse(
            "<h2>Something went wrong while generating the report.</h2>"
            "<p>Please try again. If the problem persists, contact support.</p>",
            status_code=500
        )


# =========================
# PDF REPORT
# =========================
def build_pdf_report(report_id, timestamp, patient_name, age, gender,
                     label, headline, screening_score, confidence, is_positive,
                     image_path, image_shape, vitals, lab_rows, heatmap_info,
                     region_localized, severity_info=None, stage1_result=None, urgent_override=False,
                     result_class="negative", xray_positive=False):

    pdf_path = f"reports/report_{report_id}.pdf"

    doc = SimpleDocTemplate(
        pdf_path, pagesize=A4,
        topMargin=18 * mm, bottomMargin=18 * mm, leftMargin=18 * mm, rightMargin=18 * mm
    )

    styles = getSampleStyleSheet()
    esc = html.escape

    title_style = ParagraphStyle("TitleStyle", parent=styles["Title"], fontSize=20,
                                 textColor=colors.HexColor("#0f172a"), alignment=TA_LEFT, spaceAfter=2)
    sub_style = ParagraphStyle("SubStyle", parent=styles["Normal"], fontSize=10,
                               textColor=colors.HexColor("#64748b"), alignment=TA_LEFT)
    section_style = ParagraphStyle("SectionStyle", parent=styles["Heading2"], fontSize=12.5,
                                   textColor=colors.HexColor("#0ea5e9"), spaceBefore=14, spaceAfter=8)
    normal_style = ParagraphStyle("NormalStyle", parent=styles["Normal"], fontSize=10.5, leading=15)
    headline_colors = {"positive": "#b91c1c", "caution": "#b45309", "negative": "#15803d"}
    result_style = ParagraphStyle(
        "ResultStyle", parent=styles["Normal"], fontSize=16, leading=20, alignment=TA_CENTER,
        textColor=colors.HexColor(headline_colors.get(result_class, "#15803d")),
        fontName="Helvetica-Bold")
    conf_style = ParagraphStyle("ConfStyle", parent=styles["Normal"], fontSize=10.5, alignment=TA_CENTER,
                                textColor=colors.HexColor("#475569"), leading=14)
    disclaimer_style = ParagraphStyle("DisclaimerStyle", parent=styles["Normal"], fontSize=8,
                                      textColor=colors.HexColor("#94a3b8"), alignment=TA_CENTER, leading=11)
    lab_header_style = ParagraphStyle("LabHeaderStyle", parent=styles["Normal"], fontSize=9,
                                      textColor=colors.white, fontName="Helvetica-Bold")
    lab_cell_style = ParagraphStyle("LabCellStyle", parent=styles["Normal"], fontSize=9.5,
                                    textColor=colors.HexColor("#0f172a"))
    small_style = ParagraphStyle("SmallStyle", parent=styles["Normal"], fontSize=9, leading=13,
                                 textColor=colors.HexColor("#334155"))

    content = []

    # ---- Header ----
    header_table = Table(
        [[Paragraph("Pneumonia Detection System", title_style),
          Paragraph(f"Report ID: {report_id}<br/>{timestamp}", sub_style)]],
        colWidths=[110 * mm, 60 * mm]
    )
    header_table.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("ALIGN", (1, 0), (1, 0), "RIGHT")]))
    content.append(header_table)
    content.append(Paragraph("AI-Assisted Chest X-Ray Screening Report", sub_style))
    content.append(Spacer(1, 6))
    content.append(HRFlowable(width="100%", thickness=1.2, color=colors.HexColor("#0ea5e9")))
    content.append(Spacer(1, 12))

    # ---- Patient info ----
    content.append(Paragraph("PATIENT INFORMATION", section_style))
    patient_table = Table(
        [["Name", patient_name, "Age", age], ["Gender", gender, "Report Date", timestamp]],
        colWidths=[28 * mm, 57 * mm, 28 * mm, 57 * mm]
    )
    patient_table.setStyle(TableStyle([
        ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
        ("FONTNAME", (2, 0), (2, -1), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 10),
        ("TEXTCOLOR", (0, 0), (0, -1), colors.HexColor("#64748b")),
        ("TEXTCOLOR", (2, 0), (2, -1), colors.HexColor("#64748b")),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 8), ("TOPPADDING", (0, 0), (-1, -1), 8),
        ("LINEBELOW", (0, 0), (-1, 0), 0.5, colors.HexColor("#e2e8f0")),
        ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#f8fafc")),
        ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#e2e8f0")),
        ("LEFTPADDING", (0, 0), (-1, -1), 10),
    ]))
    content.append(patient_table)
    content.append(Spacer(1, 6))

    # ---- Screening result ----
    content.append(Paragraph("SCREENING RESULT", section_style))
    box_palette = {"positive": ("#fef2f2", "#fecaca"),
                   "caution": ("#fffbeb", "#fde68a"),
                   "negative": ("#f0fdf4", "#bbf7d0")}
    bg_hex, border_hex = box_palette.get(result_class, box_palette["negative"])
    result_box_color, result_border_color = colors.HexColor(bg_hex), colors.HexColor(border_hex)
    result_inner = Table(
        [[Paragraph(headline, result_style)],
         [Paragraph(f"Pneumonia likelihood: {screening_score:.1f}%<br/>"
                    f"<font size=8>Combined likelihood from the available checks (X-ray weighted highest). "
                    f"Not a diagnosis.</font>", conf_style)]],
        colWidths=[170 * mm]
    )
    result_inner.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), result_box_color),
        ("BOX", (0, 0), (-1, -1), 1, result_border_color),
        ("TOPPADDING", (0, 0), (-1, 0), 14), ("BOTTOMPADDING", (0, 0), (-1, 0), 4),
        ("BOTTOMPADDING", (0, 1), (-1, 1), 14),
    ]))
    content.append(result_inner)
    content.append(Spacer(1, 10))

    # ---- Severity (patient wording) ----
    if severity_info is not None:
        content.append(Paragraph("SEVERITY CHECK", section_style))
        if severity_info and severity_info["available"]:
            p = severity_info["patient"]
            sev_is_severe = severity_info["level"] == "severe"
            sev_box_color = colors.HexColor("#fef2f2") if sev_is_severe else colors.HexColor("#fffbeb")
            sev_border_color = colors.HexColor("#fecaca") if sev_is_severe else colors.HexColor("#fde68a")
            sev_text_color = colors.HexColor("#b91c1c") if sev_is_severe else colors.HexColor("#b45309")
            sev_label_style = ParagraphStyle("SevLabelStyle", parent=styles["Normal"], fontSize=14, leading=18,
                                             alignment=TA_CENTER, textColor=sev_text_color, fontName="Helvetica-Bold")
            sev_desc_style = ParagraphStyle("SevDescStyle", parent=styles["Normal"], fontSize=9.5,
                                            alignment=TA_CENTER, leading=14, textColor=colors.HexColor("#334155"))
            extra = ""
            if severity_info["provisional"] and not sev_is_severe:
                extra = (" Some checks could not be done because values were not entered, "
                         "so this is only a first estimate.")
            rows = [
                [Paragraph(f"{p['title']} ({p['tag']})", sev_label_style)],
                [Paragraph(p["text"] + extra, sev_desc_style)],
            ]
            sev_inner = Table(rows, colWidths=[170 * mm])
            sev_inner.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, -1), sev_box_color),
                ("BOX", (0, 0), (-1, -1), 1, sev_border_color),
                ("TOPPADDING", (0, 0), (-1, 0), 12), ("BOTTOMPADDING", (0, 0), (-1, 0), 4),
                ("LEFTPADDING", (0, 1), (-1, -1), 16), ("RIGHTPADDING", (0, 1), (-1, -1), 16),
                ("BOTTOMPADDING", (0, 1), (-1, -1), 10),
            ]))
            content.append(sev_inner)
        else:
            reason = severity_info["reason"] if severity_info else "Severity could not be assessed."
            content.append(Paragraph(f"Severity could not be assessed: {esc(reason)}", normal_style))
        content.append(Spacer(1, 10))

    # ---- Abnormal values note + Recommendation ----
    notes = build_value_notes(is_positive, lab_rows, stage1_result, severity_info, age)
    note_heading_style = ParagraphStyle("NoteHeading", parent=styles["Heading2"], fontSize=12.5,
                                        textColor=colors.HexColor("#0f172a"), spaceBefore=10, spaceAfter=4)
    note_text_style = ParagraphStyle("NoteText", parent=styles["Normal"], fontSize=10, leading=15,
                                     textColor=colors.HexColor("#334155"))
    note_palette = {"amber": ("#fffbeb", "#fde68a", "#d97706"),
                    "red": ("#fef2f2", "#fecaca", "#dc2626"),
                    "green": ("#f0fdf4", "#bbf7d0", "#16a34a")}

    def _note_block(heading, text, tone):
        bg, border, bar = note_palette[tone]
        box = Table([[Paragraph(esc(text), note_text_style)]], colWidths=[170 * mm])
        box.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor(bg)),
            ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor(border)),
            ("LINEBEFORE", (0, 0), (0, -1), 4, colors.HexColor(bar)),
            ("TOPPADDING", (0, 0), (-1, -1), 10), ("BOTTOMPADDING", (0, 0), (-1, -1), 10),
            ("LEFTPADDING", (0, 0), (-1, -1), 14), ("RIGHTPADDING", (0, 0), (-1, -1), 12),
        ]))
        return [Paragraph(heading, note_heading_style),
                HRFlowable(width="100%", thickness=0.7, color=colors.HexColor("#cbd5e1")),
                Spacer(1, 6), box, Spacer(1, 10)]

    if notes["abnormal_text"]:
        content.extend(_note_block("PLEASE NOTE: ABNORMAL VALUES", notes["abnormal_text"], "amber"))
    content.extend(_note_block("RECOMMENDATION", notes["recommendation"], notes["tone"]))

    # ---- Summary of checks (patient wording) ----
    if stage1_result is not None:
        content.append(Paragraph("SUMMARY OF CHECKS", section_style))
        data = [[Paragraph("CHECK", lab_header_style), Paragraph("RESULT", lab_header_style),
                 Paragraph("BASIS", lab_header_style)]]
        oxygen_grade = stage1_result.get("oxygen_grade")
        blood_abnormal = stage1_result.get("blood_abnormal")
        for c in stage1_result["components"]:
            name, text, pill, note = friendly_signal(c, oxygen_grade, blood_abnormal, urgent_override)
            col = {"pos": "#b91c1c", "neg": "#15803d", "warn": "#b45309", "na": "#94a3b8"}[pill]
            rs = ParagraphStyle(f"S1R{c['name']}", parent=styles["Normal"], fontSize=9,
                                fontName="Helvetica-Bold", textColor=colors.HexColor(col))
            data.append([Paragraph(esc(name), lab_cell_style), Paragraph(esc(text), rs),
                         Paragraph(esc(note) or "-", small_style)])
        t = Table(data, colWidths=[50 * mm, 45 * mm, 75 * mm])
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0ea5e9")),
            ("TOPPADDING", (0, 0), (-1, -1), 7), ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
            ("LEFTPADDING", (0, 0), (-1, -1), 8),
            ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#e2e8f0")),
            ("LINEBELOW", (0, 0), (-1, -2), 0.5, colors.HexColor("#e2e8f0")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f8fafc")]),
        ]))
        content.append(t)
        if stage1_result["disagree"]:
            content.append(Spacer(1, 4))
            content.append(Paragraph("Some checks disagree with each other, so a doctor's review is strongly advised.",
                                     small_style))
        content.append(Spacer(1, 10))

    # ---- Vitals & labs ----
    content.append(Paragraph("VITALS &amp; LAB VALUES", section_style))
    flag_colors = {"Normal": colors.HexColor("#15803d"), "Low": colors.HexColor("#b45309"),
                   "High": colors.HexColor("#b91c1c"), "N/A": colors.HexColor("#94a3b8")}
    lab_table_data = [[Paragraph(h, lab_header_style) for h in ("TEST", "RESULT", "UNIT", "NORMAL RANGE", "STATUS")]]
    for i, row in enumerate(lab_rows):
        flag_style = ParagraphStyle(f"FlagStyle{i}", parent=styles["Normal"], fontSize=9.5,
                                    fontName="Helvetica-Bold",
                                    textColor=flag_colors.get(row["flag"], colors.HexColor("#0f172a")))
        lab_table_data.append([
            Paragraph(row["label"], lab_cell_style), Paragraph(row["value"], lab_cell_style),
            Paragraph(row["unit"], lab_cell_style), Paragraph(row["range"], lab_cell_style),
            Paragraph(row["flag"], flag_style),
        ])
    lab_table = Table(lab_table_data, colWidths=[38 * mm, 28 * mm, 28 * mm, 40 * mm, 36 * mm])
    lab_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0ea5e9")),
        ("TOPPADDING", (0, 0), (-1, 0), 8), ("BOTTOMPADDING", (0, 0), (-1, 0), 8),
        ("TOPPADDING", (0, 1), (-1, -1), 7), ("BOTTOMPADDING", (0, 1), (-1, -1), 7),
        ("LEFTPADDING", (0, 0), (-1, -1), 8),
        ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#e2e8f0")),
        ("LINEBELOW", (0, 0), (-1, -2), 0.5, colors.HexColor("#e2e8f0")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f8fafc")]),
    ]))
    content.append(lab_table)
    content.append(Spacer(1, 4))
    content.append(Paragraph(
        "Reference ranges are standard adult values and may vary by laboratory and patient.",
        ParagraphStyle("RangeNote", parent=styles["Normal"], fontSize=7.5, textColor=colors.HexColor("#94a3b8"))))
    content.append(Spacer(1, 10))

    # ---- X-ray ----
    xray_block = [Paragraph("CHEST X-RAY", section_style)]
    try:
        px_h, px_w = image_shape[0], image_shape[1]
        target_w = 170 * mm
        target_h = target_w * (px_h / px_w)
        img_table = Table([[Image(image_path, width=target_w, height=target_h)]], colWidths=[170 * mm])
        img_table.setStyle(TableStyle([
            ("ALIGN", (0, 0), (-1, -1), "CENTER"),
            ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#0f172a")),
            ("TOPPADDING", (0, 0), (-1, -1), 8), ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
        ]))
        xray_block.append(img_table)
        xray_block.append(Spacer(1, 4))
        if xray_positive:
            cap = ("Yellow box marks the area of concern; the red/yellow heatmap shows where the AI focused."
                   if region_localized else
                   "The red/yellow heatmap shows where the AI focused. No single area was distinct enough to box.")
        else:
            cap = "Original chest X-ray (heatmap hidden because the X-ray model found no abnormal region)."
        xray_block.append(Paragraph(cap, conf_style))
    except Exception:
        xray_block.append(Paragraph("X-ray image could not be embedded.", normal_style))
    content.append(KeepTogether(xray_block))
    content.append(Spacer(1, 16))

    # ---- Signature ----
    sign_label_style = ParagraphStyle("SignLabelStyle", parent=styles["Normal"], fontSize=8.5,
                                      textColor=colors.HexColor("#94a3b8"))
    sign_table = Table(
        [["", ""], [Paragraph("Authorized Signatory", sign_label_style), Paragraph("Date", sign_label_style)]],
        colWidths=[85 * mm, 85 * mm], rowHeights=[16 * mm, 5 * mm]
    )
    sign_table.setStyle(TableStyle([
        ("LINEABOVE", (0, 1), (0, 1), 0.7, colors.HexColor("#cbd5e1")),
        ("LINEABOVE", (1, 1), (1, 1), 0.7, colors.HexColor("#cbd5e1")),
        ("TOPPADDING", (0, 1), (-1, 1), 4),
    ]))
    content.append(KeepTogether([sign_table]))
    content.append(Spacer(1, 12))
    content.append(HRFlowable(width="100%", thickness=0.7, color=colors.HexColor("#e2e8f0")))
    content.append(Spacer(1, 8))
    content.append(Paragraph(SCREENING_DISCLAIMER, disclaimer_style))

    # =====================================================
    # CLINICAL APPENDIX (for doctors) - separate last page
    # =====================================================
    content.append(PageBreak())
    content.append(Paragraph("CLINICAL DETAILS (FOR DOCTORS)", section_style))

    if stage1_result is not None:
        content.append(Paragraph(
            f"Decision rule: the X-ray model is the primary signal; a positive X-ray is never overridden by "
            f"the other models. Combined pneumonia likelihood (weighted, X-ray highest) = "
            f"{stage1_result['final_prob'] * 100:.1f}% &rarr; <b>{esc(stage1_result['display_label'])}</b>"
            + (" (signals disagree)" if stage1_result["disagree"] else ""), small_style))
        if stage1_result.get("reasons"):
            content.append(Spacer(1, 3))
            content.append(Paragraph("<br/>".join(f"&bull; {esc(r)}" for r in stage1_result["reasons"]),
                                     small_style))
        content.append(Spacer(1, 6))
        data = [[Paragraph(h, lab_header_style) for h in ("SIGNAL", "RESULT", "DETAIL")]]
        tech_blood_abnormal = stage1_result.get("blood_abnormal")
        for c in stage1_result["components"]:
            pill, pill_text, detail = technical_signal(c, tech_blood_abnormal, stage1_result.get("oxygen_grade"))
            col = {"pos": "#b91c1c", "neg": "#15803d", "warn": "#b45309", "na": "#94a3b8"}[pill]
            rs = ParagraphStyle(f"TechR{c['name']}", parent=styles["Normal"], fontSize=9,
                                fontName="Helvetica-Bold", textColor=colors.HexColor(col))
            data.append([Paragraph(esc(c["name"]), small_style), Paragraph(esc(pill_text), rs),
                         Paragraph(esc(detail) or "-", small_style)])
        t = Table(data, colWidths=[40 * mm, 40 * mm, 90 * mm])
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0ea5e9")),
            ("TOPPADDING", (0, 0), (-1, -1), 6), ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
            ("LEFTPADDING", (0, 0), (-1, -1), 8),
            ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#e2e8f0")),
            ("LINEBELOW", (0, 0), (-1, -2), 0.5, colors.HexColor("#e2e8f0")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f8fafc")]),
        ]))
        content.append(t)
        content.append(Spacer(1, 6))

        og = stage1_result.get("oxygen_grade")
        if og and og.get("available"):
            content.append(Paragraph(
                f"<b>Oxygen &amp; pulse grading</b> ({esc(og['situation_label'])} SpO2 thresholds): "
                f"{esc(og['breakdown'])} &rarr; <b>{esc(og['label'])}</b> (worst value decides).", small_style))
        if urgent_override:
            content.append(Spacer(1, 4))
            content.append(Paragraph(
                "<b>Safety rule applied:</b> no pneumonia pattern was found, but the Oxygen &amp; pulse "
                "grade is Abnormal, so the headline shows urgent medical review advised instead of an "
                "unqualified \"no signs\" message.", small_style))
        content.append(Spacer(1, 10))

    if is_positive and severity_info and severity_info["available"]:
        prov = " (provisional)" if severity_info["provisional"] else ""
        content.append(Paragraph(
            f"<b>Severity staging: {severity_info['label']}{prov}</b> &mdash; {esc(severity_info['full_name'])}. "
            f"{severity_info['met']} of {severity_info['evaluated']} evaluable ATS/IDSA 2007 minor criteria met "
            f"(SCAP needs 3).", small_style))
        content.append(Spacer(1, 4))
        lines = "<br/>".join(
            f"&bull; {esc(b['name'])}: {esc(b['status'])} ({esc(b['detail'])})" for b in severity_info["criteria"])
        content.append(Paragraph(lines, small_style))
        if severity_info.get("urgent_override") and severity_info["label"] != "SCAP":
            content.append(Spacer(1, 4))
            content.append(Paragraph(
                "<b>Safety rule applied:</b> an Abnormal (clearly low) SpO2 for this patient's situation "
                "triggers the urgent message regardless of the criteria count.", small_style))
        content.append(Spacer(1, 4))
        content.append(Paragraph(esc(severity_info["description"]), small_style))
        content.append(Spacer(1, 10))

    content.append(Paragraph("Grad-CAM interpretation", ParagraphStyle(
        "GH", parent=styles["Normal"], fontSize=10, fontName="Helvetica-Bold",
        textColor=colors.HexColor("#0f172a"))))
    content.append(Spacer(1, 3))
    content.append(Paragraph(
        f"Attention intensity: <b>{esc(heatmap_info['intensity'])}</b> &nbsp;|&nbsp; "
        f"Region of focus: <b>{esc(heatmap_info['location'])}</b><br/>{esc(heatmap_info['note'])}", small_style))
    content.append(Spacer(1, 10))

    # CAP / SCAP legend only when pneumonia is suspected
    if severity_info is not None:
        content.append(Paragraph("About severity staging (NSCAP / SCAP)", ParagraphStyle(
            "GH2", parent=styles["Normal"], fontSize=10, fontName="Helvetica-Bold",
            textColor=colors.HexColor("#0f172a"))))
        content.append(Spacer(1, 3))
        content.append(Paragraph(GLOSSARY_INTRO_TEXT, small_style))
        for key in ("NSCAP", "SCAP"):
            meta = SEVERITY_INFO[key]
            content.append(Spacer(1, 4))
            content.append(Paragraph(f"<b>{key}</b> &mdash; {meta['full_name']}: {meta['description']}", small_style))

    def add_page_footer(canvas_obj, doc_obj):
        canvas_obj.saveState()
        canvas_obj.setFont("Helvetica", 8)
        canvas_obj.setFillColor(colors.HexColor("#94a3b8"))
        canvas_obj.drawCentredString(
            A4[0] / 2, 12 * mm,
            f"Pneumonia Detection System  |  Report #{report_id}  |  Page {doc_obj.page}")
        canvas_obj.restoreState()

    doc.build(content, onFirstPage=add_page_footer, onLaterPages=add_page_footer)
    return pdf_path


# =========================
# HISTORY
# =========================
@app.get("/history", response_class=HTMLResponse)
def history():
    try:
        conn = get_connection()
        try:
            cursor = conn.cursor(dictionary=True)
            try:
                cursor.execute("""
                    SELECT dr.report_id, dr.report_code, dr.final_diagnosis,
                           dr.confidence, dr.image_name, dr.created_at, dr.severity,
                           p.full_name, p.age, p.gender
                    FROM diagnosis_reports dr
                    JOIN patients p ON dr.patient_id = p.patient_id
                    ORDER BY dr.report_id DESC
                """)
                rows = cursor.fetchall()
            finally:
                cursor.close()
        finally:
            conn.close()
        db_error = None
    except Exception as e:
        rows = []
        db_error = str(e)

    if db_error:
        rows_html = f"""
        <tr><td colspan="9" style="text-align:center;padding:24px;color:#b91c1c;">
            Could not load records: {html.escape(db_error)}
        </td></tr>"""
    elif not rows:
        rows_html = """
        <tr><td colspan="9" style="text-align:center;padding:24px;color:#94a3b8;">
            No records saved yet. Run a scan from the home page first.
        </td></tr>"""
    else:
        rows_html = ""
        severity_names = {"NSCAP": "Mild to Moderate", "SCAP": "Severe"}
        for r in rows:
            diag = (r["final_diagnosis"] or "").upper()
            if "INCONCLUSIVE" in diag:
                flag_class, result_text = "Low", "Inconclusive"
            elif "POSSIBLE" in diag:
                flag_class, result_text = "Low", "Possible pneumonia"
            elif "PNEUMONIA" in diag:
                flag_class, result_text = "High", "Pneumonia suspected"
            else:
                flag_class, result_text = "Normal", "No signs found"
            created = r["created_at"].strftime("%d %b %Y, %I:%M %p") if r["created_at"] else "-"
            sev_raw = r.get("severity")
            severity_val = severity_names.get(sev_raw, "-") if sev_raw else "-"
            rows_html += f"""
            <tr>
                <td>{html.escape(r['report_code'] or '-')}</td>
                <td>{html.escape(r['full_name'] or '-')}</td>
                <td>{html.escape(str(r['age']) if r['age'] is not None else '-')}</td>
                <td>{html.escape(r['gender'] or '-')}</td>
                <td><span class="flag-pill {flag_class}">{result_text}</span></td>
                <td>{r['confidence']}%</td>
                <td class="range-col">{html.escape(severity_val)}</td>
                <td class="range-col">{created}</td>
                <td><a href="/download?report_id={r['report_code']}" class="btn btn-outline" style="width:auto;padding:6px 14px;font-size:12.5px;display:inline-flex;">⬇ PDF</a></td>
            </tr>"""

    return f"""
<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>Saved Records &middot; Pneumonia Detection System</title>
<style>{BASE_CSS}</style>
</head>
<body>

<div class="shell">

  <div class="brand">
    <div class="logo">🫁</div>
    <h1>Pneumonia Detection System
      <span>All Saved Patients &amp; Reports</span>
    </h1>
  </div>

  <div class="card">
    <div class="card-pad">

      <div class="actions" style="margin-bottom:22px;">
        <a href="/" class="btn btn-outline">↺ New Scan</a>
        <a href="{STREAMLIT_URL}" class="btn btn-outline">Go to Home</a>
      </div>

      <div class="lab-card">
        <table class="lab-table">
          <thead>
            <tr>
              <th>Report ID</th><th>Patient</th><th>Age</th><th>Gender</th>
              <th>Result</th><th>Likelihood</th><th>Severity</th><th>Date</th><th>PDF</th>
            </tr>
          </thead>
          <tbody>{rows_html}</tbody>
        </table>
      </div>

    </div>
  </div>

</div>

</body>
</html>
"""


# =========================
# PDF DOWNLOAD
# =========================
def find_pdf_path(report_id: str):
    report_id = report_id.strip().upper()

    pdf_path = PDF_STORE.get(report_id)
    if pdf_path and os.path.isfile(pdf_path):
        return os.path.abspath(pdf_path)

    reports_folder = os.path.join(os.path.dirname(os.path.abspath(__file__)), "reports")

    try:
        conn = get_connection()
        try:
            cursor = conn.cursor(dictionary=True)
            try:
                cursor.execute(
                    "SELECT pdf_path FROM diagnosis_reports WHERE UPPER(report_code) = %s LIMIT 1",
                    (report_id,)
                )
                row = cursor.fetchone()
            finally:
                cursor.close()
        finally:
            conn.close()

        if row and row.get("pdf_path"):
            db_path = os.path.normpath(row["pdf_path"])
            if os.path.isfile(db_path):
                PDF_STORE[report_id] = db_path
                return os.path.abspath(db_path)

            local_pdf_path = os.path.join(reports_folder, os.path.basename(db_path))
            if os.path.isfile(local_pdf_path):
                PDF_STORE[report_id] = local_pdf_path
                return os.path.abspath(local_pdf_path)

    except Exception as e:
        print("Database PDF lookup error:", e)

    fallback_path = os.path.join(reports_folder, f"report_{report_id}.pdf")
    if os.path.isfile(fallback_path):
        PDF_STORE[report_id] = fallback_path
        return os.path.abspath(fallback_path)

    print("PDF NOT FOUND FOR REPORT:", report_id)
    return None


def _serve_pdf(report_id: str):
    if not report_id:
        raise HTTPException(status_code=400, detail="Report ID is required")

    pdf_path = find_pdf_path(report_id)
    if not pdf_path:
        raise HTTPException(status_code=404, detail=f"PDF report not found for report ID: {report_id}")

    return FileResponse(
        path=pdf_path,
        media_type="application/pdf",
        filename=os.path.basename(pdf_path),
        content_disposition_type="inline"
    )


@app.get("/download")
def download_pdf(report_id: str = ""):
    return _serve_pdf(report_id)


@app.get("/download/{report_id}")
def view_pdf(report_id: str):
    return _serve_pdf(report_id)


# =========================
# JSON ENDPOINTS
# =========================
@app.get("/patients")
def get_patients():
    conn = get_connection()
    try:
        cursor = conn.cursor(dictionary=True)
        try:
            cursor.execute("SELECT * FROM patients ORDER BY patient_id DESC")
            return cursor.fetchall()
        finally:
            cursor.close()
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))
    finally:
        conn.close()


@app.get("/reports")
def get_reports():
    conn = get_connection()
    try:
        cursor = conn.cursor(dictionary=True)
        try:
            cursor.execute("SELECT * FROM diagnosis_reports ORDER BY report_id DESC")
            return cursor.fetchall()
        finally:
            cursor.close()
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))
    finally:
        conn.close()